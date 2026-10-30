"""Recrea las dos fotografías como tablas editables y PDF vectorial.
Ejecutar: python images/convertidos/generar.py
"""
import json
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from xml.sax.saxutils import escape
from PIL import Image, ImageOps
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.worksheet.page import PageMargins
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.colors import HexColor
from reportlab.platypus import Table, TableStyle, Paragraph
from reportlab.lib.styles import ParagraphStyle

HERE = Path(__file__).resolve().parent
DATA = json.loads((HERE / 'datos.json').read_text(encoding='utf-8'))
TITLE = DATA['titulo']
YELLOW, PINK, MINT = 'FFF176', 'F5AEC8', 'D9EEE5'
HEAD = ['Descripción Cargo', 'Cat.\nOcup.', 'Cant.\nCargos', 'Escala', 'Salario', 'Cant.', 'Escala', 'Salario', 'Observación']
WIDTHS = [245, 45, 54, 52, 72, 47, 57, 72, 135]

def make_sheet(wb, key):
    rows = DATA['paginas'][key]
    ws = wb.create_sheet(key)
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    ws.merge_cells('A1:I1')
    ws['A1'] = TITLE if key == '58622' else TITLE + ' · Continuación'
    ws['A1'].font = Font(name='Calibri', size=17, bold=True)
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 35
    for area, value in [('A2:A3', HEAD[0]), ('B2:B3', HEAD[1]), ('C2:C3', HEAD[2]), ('D2:E2', 'Actual'), ('F2:H2', 'Propuesta'), ('I2:I3', HEAD[8])]:
        ws.merge_cells(area)
        cell = ws[area.split(':')[0]]
        cell.value = value
    for col in range(4, 9):
        ws.cell(3, col, HEAD[col-1])
    line = Side(style='thin', color='555555')
    for row in ws.iter_rows(min_row=2, max_row=3, max_col=9):
        for cell in row:
            cell.font = Font(name='Calibri', size=12, bold=True)
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            cell.border = Border(left=line, right=line, top=line, bottom=line)
    for addr in ['D2','F2']:
        ws[addr].fill = PatternFill('solid', fgColor=YELLOW)
    for addr in ['C2','F3']:
        ws[addr].fill = PatternFill('solid', fgColor=PINK)
    ws['E3'].fill = PatternFill('solid', fgColor=MINT)
    ws.row_dimensions[2].height = 24
    ws.row_dimensions[3].height = 28
    for r, values in enumerate(rows, 4):
        ws.row_dimensions[r].height = 26
        for c, value in enumerate(values, 1):
            cell = ws.cell(r, c, value)
            cell.font = Font(name='Calibri', size=12)
            cell.alignment = Alignment(horizontal='left' if c in (1,9) else 'center', vertical='center', wrap_text=True)
            cell.border = Border(left=line, right=line, top=line, bottom=line)
            if c in (5,8) and isinstance(value,(int,float)):
                cell.number_format = '0.00'
    for col,width in zip('ABCDEFGHI',[43,8,10,10,13,9,11,13,23]):
        ws.column_dimensions[col].width = width
    for label, note in DATA['notas']:
        if label.startswith(key) and ' · ' in label:
            addr = label.split(' · ')[-1]
            ws[addr].comment = Comment(note, 'Nota de transcripción')
    last = 3 + len(rows)
    ws.freeze_panes = 'D4'
    ws.print_area = f'A1:I{last}'
    ws.print_title_rows = '1:3'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = ws.page_setup.fitToHeight = 1
    ws.page_margins = PageMargins(left=.25,right=.25,top=.35,bottom=.35,header=.15,footer=.15)
    ws.oddFooter.center.text = f'Fuente: images/{key}.jpg · Recreación tipográfica · Consulte Notas'
    ws.oddFooter.center.size = 8
    return ws

def workbook(keys, target):
    wb = Workbook()
    wb.remove(wb.active)
    for key in keys:
        make_sheet(wb,key)
    notes = wb.create_sheet('Notas')
    notes.append(['Referencia','Criterio de transcripción / lectura a confirmar'])
    for row in DATA['notas']:
        notes.append(row)
    notes.column_dimensions['A'].width = 30
    notes.column_dimensions['B'].width = 115
    for row in notes:
        for cell in row:
            cell.alignment = Alignment(vertical='top',wrap_text=True)
            cell.font = Font(name='Calibri',size=11,bold=cell.row==1)
        notes.row_dimensions[row[0].row].height = 49 if row[0].row > 1 else 30
    notes.freeze_panes = 'B2'
    for key in keys:
        ws = wb.create_sheet(f'Foto {key}')
        ws['A1'] = f'Referencia original: images/{key}.jpg (incluye rúbrica original si existe)'
        ws['A1'].font = Font(bold=True,size=13)
        im = ImageOps.exif_transpose(Image.open(HERE.parent / f'{key}.jpg'))
        buf = BytesIO()
        im.save(buf,format='JPEG',quality=93)
        buf.seek(0)
        image = ExcelImage(buf)
        image.width, image.height = 1248,936
        ws.add_image(image,'A3')
        ws.sheet_view.showGridLines = False
    wb.save(target)

STYLE = ParagraphStyle('cell',fontName='Helvetica',fontSize=10,leading=12)
BOLD = ParagraphStyle('head',parent=STYLE,fontName='Helvetica-Bold',alignment=1)

def draw_page(c,key):
    w,h = landscape(A4)
    margin=24
    available=w-2*margin
    widths=[v*available/sum(WIDTHS) for v in WIDTHS]
    c.setTitle(TITLE)
    c.setAuthor('Transcripción de fotografías del repositorio Capital_Humanos')
    c.setFont('Helvetica-Bold',14)
    c.drawCentredString(w/2,h-36,TITLE)
    c.setFont('Helvetica',9)
    c.drawCentredString(w/2,h-51,'Página 1 · 58622.jpg' if key=='58622' else 'Página 2 · 58623.jpg · Continuación')
    p=lambda s:Paragraph(escape(str(s)).replace('\n','<br/>'),STYLE)
    ph=lambda s:Paragraph(escape(s).replace('\n','<br/>'),BOLD)
    table_rows=[list(map(ph,[HEAD[0],HEAD[1],HEAD[2],'Actual','','Propuesta','','',HEAD[8]])),
                [ph(''),ph(''),ph(''),ph('Escala'),ph('Salario'),ph('Cant.'),ph('Escala'),ph('Salario'),ph('')]]
    for values in DATA['paginas'][key]:
        display=[f'{v:.2f}' if j in (4,7) and isinstance(v,(float,int)) else v for j,v in enumerate(values)]
        table_rows.append([p(v) for v in display])
    # Same row height in both pages, preserving a blank lower area on page 2.
    table=Table(table_rows,colWidths=widths,rowHeights=[22,25]+[19]*len(DATA['paginas'][key]))
    cmds=[('GRID',(0,0),(-1,-1),.45,HexColor('#555555')),('VALIGN',(0,0),(-1,-1),'MIDDLE'),
          ('SPAN',(0,0),(0,1)),('SPAN',(1,0),(1,1)),('SPAN',(2,0),(2,1)),
          ('SPAN',(3,0),(4,0)),('SPAN',(5,0),(7,0)),('SPAN',(8,0),(8,1)),
          ('BACKGROUND',(3,0),(4,0),HexColor('#'+YELLOW)),('BACKGROUND',(5,0),(7,0),HexColor('#'+YELLOW)),
          ('BACKGROUND',(2,0),(2,1),HexColor('#'+PINK)),('BACKGROUND',(5,1),(5,1),HexColor('#'+PINK)),
          ('BACKGROUND',(4,1),(4,1),HexColor('#'+MINT)),('LEFTPADDING',(0,0),(-1,-1),5),('RIGHTPADDING',(0,0),(-1,-1),5)]
    # Paragraph alignment is set separately to keep all numeric columns centered.
    for row in table_rows[2:]:
        for j in range(1,8):
            row[j].style=ParagraphStyle('center',parent=STYLE,alignment=1)
    table.setStyle(TableStyle(cmds))
    tw,th=table.wrap(available,h)
    table.drawOn(c,margin,h-66-th)
    c.setFont('Helvetica',8)
    c.drawString(margin,39,'Recreación tipográfica. Se conservan los valores manuscritos; no constituye validación salarial.')
    c.drawString(margin,27,'Lecturas a confirmar y criterios: véase Notas de transcripción. La rúbrica no se reproduce como firma.')
    c.showPage()

def pdf(keys,target):
    c=canvas.Canvas(str(target),pagesize=landscape(A4),pageCompression=1)
    for key in keys:
        draw_page(c,key)
    c.save()

def main():
    for key in DATA['paginas']:
        workbook([key],HERE/f'{key}.xlsx')
        pdf([key],HERE/f'{key}.pdf')
    workbook(list(DATA['paginas']),HERE/'Plantilla_Campismo.xlsx')
    pdf(list(DATA['paginas']),HERE/'Plantilla_Campismo.pdf')
    lines=['# Notas de transcripción','', 'Los números de fila se refieren al orden de los cargos en cada fotografía; las direcciones de celda corresponden al Excel.','']
    for label,note in DATA['notas']:
        lines.extend([f'## {label}',note,''])
    (HERE/'NOTAS.md').write_text('\n'.join(lines),encoding='utf-8')
    with ZipFile(HERE/'Conversion_58622_58623.zip','w',ZIP_DEFLATED) as z:
        for name in ['58622.xlsx','58622.pdf','58623.xlsx','58623.pdf','Plantilla_Campismo.xlsx','Plantilla_Campismo.pdf','NOTAS.md','README.md']:
            z.write(HERE/name,arcname=name)
    print('Generados: 3 Excel, 3 PDF y paquete ZIP.')

if __name__=='__main__':
    main()
