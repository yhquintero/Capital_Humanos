from __future__ import annotations

import contextlib
import hashlib
import io
import json
import math
import runpy
import sys
import types
import zipfile
from io import BytesIO
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
APP_DIR = ROOT / "TES" / "Salariales" / "Nueva Escala 3210"
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(APP_DIR))

import generar_escala as payroll  # noqa: E402
from tools.verificar_paquete_exportacion import verificar_paquete  # noqa: E402


@pytest.fixture(scope="module")
def app_namespace():
    # Load export helpers without starting a Streamlit server.
    with contextlib.redirect_stderr(io.StringIO()):
        return runpy.run_path(str(APP_DIR / "app_escala.py"), run_name="capital_humano_export_test")


def test_app_loads_its_sibling_generator_even_if_module_name_is_shadowed(monkeypatch):
    # Reproduce a stale/foreign `generar_escala` already loaded by a Windows
    # IDE, Streamlit session, or another checkout on sys.path.
    monkeypatch.setitem(sys.modules, "generar_escala", types.ModuleType("generar_escala"))
    with contextlib.redirect_stderr(io.StringIO()):
        namespace = runpy.run_path(str(APP_DIR / "app_escala.py"), run_name="capital_humano_import_test")

    generator = namespace["base"]
    assert Path(generator.__file__).resolve() == (APP_DIR / "generar_escala.py").resolve()
    assert generator.FACTOR_MAXIMO_ABSOLUTO == 6.0


@pytest.mark.parametrize(
    ("salary", "factor", "expected"),
    [(3210, 1, 3210), (3210, 2.5, 8025), (500, 6, 3000), (500, 7, None)],
)
def test_salary_multiplier_never_exceeds_six_times(salary, factor, expected):
    if expected is None:
        with pytest.raises(ValueError):
            payroll.aplicar_factor_con_tope(salary, factor)
        return
    result = payroll.aplicar_factor_con_tope(salary, factor)
    assert result == pytest.approx(expected)
    assert result <= salary * 6


@pytest.mark.parametrize(
    ("salary", "factor"),
    [(-1, 1), (100, -0.1), (100, math.inf), (math.nan, 2), (math.inf, 1)],
)
def test_salary_multiplier_rejects_invalid_inputs(salary, factor):
    with pytest.raises(ValueError):
        payroll.aplicar_factor_con_tope(salary, factor)


def test_official_scale_has_32_groups_and_preserves_minimum():
    assert len(payroll.ESCALA_OFICIAL_RES15) == 32
    assert set(payroll.ESCALA_OFICIAL_RES15) == set(payroll.GRUPOS_ROMANOS)
    assert payroll.ESCALA_OFICIAL_RES15["I"] == 3210


def test_source_files_are_present_and_template_is_auditable():
    assert payroll.RUTA_PLANTILLA.is_file()
    assert payroll.RUTA_ESCALA_BASE.is_file()
    proposal, counters, records, anomalies = payroll.extraer_propuesta_plantilla()
    assert records
    assert len(records) >= 150
    assert set(proposal).issubset(payroll.ROMANO_A_ENTERO)
    assert all(record["hoja"] and record["fila"] > 0 and record["cargo"] for record in records)
    assert set(counters) == set(proposal)
    assert isinstance(anomalies, list)


def test_user_authorized_majority_salary_corrections_resolve_template_findings():
    proposal, counters, records, anomalies, findings = payroll.extraer_propuesta_plantilla_con_calidad()
    source = payroll.xlrd.open_workbook(str(payroll.RUTA_PLANTILLA), formatting_info=True)
    corrections = [
        ("P-2  Aguas Claras", 44, "Jefes de Grupo de Seguridad Interna", "XII"),
        ("Plantilla Emp", 51, "Director Comercial", "XV"),
    ]
    by_location = {(record["hoja"], record["fila"]): record for record in records}
    for sheet_name, row_number, cargo, group in corrections:
        record = by_location[(sheet_name, row_number)]
        sheet = source.sheet_by_name(sheet_name)
        _, group_column = payroll._localizar_columna_ge(sheet)
        source_salary = sheet.cell_value(row_number - 1, group_column + 1)
        assert record["cargo"] == cargo
        assert record["grupo"] == group
        assert source_salary == counters[group].most_common(1)[0][0]
        assert record["salario"] == source_salary
    assert len(records) == 183
    assert not anomalies
    assert not findings
    assert proposal["XII"] == counters["XII"].most_common(1)[0][0]
    assert proposal["XV"] == counters["XV"].most_common(1)[0][0]


def test_source_quality_report_flags_rows_without_mutating_or_dropping_valid_duplicates(monkeypatch):
    class FakeSheet:
        name = "Centro de prueba"
        rows = [
            ["Cargo", "G.E.A", "Salario"],
            ["Técnico de primera", "XII", 5000],
            ["Técnico de primera", "XII", 5000],
            ["Auxiliar de archivo", "XIII", ""],
            ["Supervisor de turno", "XIV", -10],
            ["Planner de mantenimiento", "XVI", "n/a"],
            ["Operador de equipo", "XII", 4800],
            ["", "XVII", 4000],
        ]
        nrows = len(rows)
        ncols = len(rows[0])

        def cell_value(self, row, column):
            return self.rows[row][column]

    class FakeWorkbook:
        def sheets(self):
            return [FakeSheet()]

    source_rows = [list(row) for row in FakeSheet.rows]
    monkeypatch.setattr(payroll.xlrd, "open_workbook", lambda _: FakeWorkbook())
    monkeypatch.setattr(payroll, "RUTA_PLANTILLA", Path("plantilla_prueba.xls"))

    proposal, counters, records, anomalies, findings = payroll.extraer_propuesta_plantilla_con_calidad()
    findings_by_type = {}
    for finding in findings:
        findings_by_type.setdefault(finding["Tipo de incidencia"], []).append(finding)

    assert proposal == {"XII": 5000.0, "XVII": 4000.0}
    assert counters["XII"] == {5000.0: 2, 4800.0: 1}
    assert [record["fila"] for record in records] == [2, 3, 7, 8]
    assert "XII" in anomalies[0]
    assert {"Salario ausente", "Salario inválido", "Cargo sin descripción", "Posible duplicado", "Discrepancia salarial de grupo"} <= set(findings_by_type)
    assert findings_by_type["Salario ausente"][0]["Fila de origen"] == 4
    assert [item["Fila de origen"] for item in findings_by_type["Salario inválido"]] == [5, 6]
    assert findings_by_type["Cargo sin descripción"][0]["Fila de origen"] == 8
    assert findings_by_type["Posible duplicado"][0]["Fila de origen"] == 3
    assert findings_by_type["Discrepancia salarial de grupo"][0]["Fila de origen"] == 7
    assert len(records) == 4  # El posible duplicado permanece; el diagnóstico no elimina datos.
    assert FakeSheet.rows == source_rows


def test_comparison_keeps_missing_groups_official_and_caps_proposed_groups():
    proposal, _, _, _ = payroll.extraer_propuesta_plantilla()
    previous = payroll.cargar_salarios_res29()
    comparison = payroll.comparar_escenarios(1, 6, proposal, previous)
    by_group = {row["Grupo"]: row for row in comparison}

    assert len(comparison) == 32
    assert by_group["I"]["Escenario A (CUP)"] == 3210
    assert by_group["I"]["Escenario B (CUP)"] == 3210
    assert by_group["XIX"]["Escenario A (CUP)"] == payroll.ESCALA_OFICIAL_RES15["XIX"]
    assert by_group["XIX"]["Escenario B (CUP)"] == payroll.ESCALA_OFICIAL_RES15["XIX"]

    proposed_group = next(iter(proposal))
    item = by_group[proposed_group]
    assert item["Escenario B (CUP)"] <= proposal[proposed_group] * 6
    assert item["Escenario B (CUP)"] == pytest.approx(proposal[proposed_group] * 6)
    assert item["Variación relativa"] == pytest.approx(
        item["Escenario B (CUP)"] / item["Escenario A (CUP)"] - 1
    )


def test_cargo_comparison_preserves_ids_origin_and_six_x_cap(app_namespace):
    _, _, records, _ = payroll.extraer_propuesta_plantilla()
    comparison = app_namespace["crear_comparacion_cargos"](1.25, 6.0, records)

    assert len(comparison) == len(records) == 183
    assert list(comparison["Id."]) == list(range(1, len(records) + 1))
    first = comparison.iloc[0]
    assert first["Hoja de origen"] == records[0]["hoja"]
    assert first["Fila de origen"] == records[0]["fila"]
    assert first["Salario A (CUP)"] == pytest.approx(first["Salario fuente (CUP)"] * 1.25)
    assert (comparison["Salario A (CUP)"] <= comparison["Salario fuente (CUP)"] * 6).all()
    assert (comparison["Salario B (CUP)"] <= comparison["Salario fuente (CUP)"] * 6).all()
    expected_relative = comparison["Salario B (CUP)"] / comparison["Salario A (CUP)"] - 1
    assert comparison["Variación relativa"].to_numpy() == pytest.approx(expected_relative.to_numpy())
    with pytest.raises(ValueError):
        app_namespace["crear_comparacion_cargos"](0.9, 1.5, records)


def test_cargo_comparison_workbook_is_readable_and_traceable(app_namespace):
    _, _, records, _ = payroll.extraer_propuesta_plantilla()
    comparison = app_namespace["crear_comparacion_cargos"](1.0, 1.5, records)
    manifest = payroll.verificar_manifiesto_fuentes()
    data = app_namespace["crear_excel_comparacion_cargos"](comparison, 1.0, 1.5, manifest)
    workbook = load_workbook(BytesIO(data), data_only=True)
    sheet = workbook["Comparación por cargo"]

    assert sheet.max_row == len(records) + 3
    assert sheet.max_column == 12
    assert sheet["A1"].value == "COMPARACIÓN DE ESCENARIOS · DETALLE POR CARGO"
    assert sheet["A4"].value == 1
    assert sheet["G4"].number_format == '0.00"×"'
    assert sheet["L4"].number_format == "0.00%"
    assert "Trazabilidad SHA-256" in workbook.sheetnames


def test_comparison_workbook_is_readable_and_contains_source_hashes(app_namespace):
    proposal, _, _, _ = payroll.extraer_propuesta_plantilla()
    previous = payroll.cargar_salarios_res29()
    comparison_rows = payroll.comparar_escenarios(1, 1.5, proposal, previous)
    frame = pd.DataFrame(comparison_rows)
    manifest = payroll.verificar_manifiesto_fuentes()
    data = app_namespace["crear_excel_comparacion"](frame, 1, 1.5, manifest)
    workbook = load_workbook(BytesIO(data), data_only=True)
    sheet = workbook["Comparación por grupo"]
    provenance = workbook["Trazabilidad SHA-256"]

    assert sheet.max_row == 35
    assert sheet.max_column == 6
    assert sheet["A1"].value == "COMPARACIÓN DE ESCENARIOS · CAPITAL HUMANO"
    assert sheet["E4"].number_format == "0.00%"
    assert sheet.freeze_panes == "A4"
    assert sheet.page_setup.orientation == "landscape"
    assert provenance["B2"].value == payroll.VERSION_APP
    assert provenance["B3"].value == "VERIFICADA"
    assert provenance["D7"].value == "VERIFICADA"
    assert len(provenance["E7"].value) == 64


def test_filtered_csv_package_contains_json_manifest(app_namespace):
    manifest = payroll.verificar_manifiesto_fuentes()
    df = pd.DataFrame([{"Id.": 1, "Hoja de origen": "Centro 1", "Cargo": "Técnico", "Salario": 3210}])
    data = app_namespace["crear_zip_csv_con_manifiesto"](df, "filtro.csv", manifest)
    with zipfile.ZipFile(BytesIO(data)) as archive:
        assert set(archive.namelist()) == {"filtro.csv", "manifiesto_fuentes.json", "LEEME.txt"}
        payload = json.loads(archive.read("manifiesto_fuentes.json"))
        assert payload["integrity_ok"] is True
        assert len(payload["sources"]) == len(payroll.FUENTES_CONTROLADAS)
        assert b"Id.,Hoja de origen,Cargo,Salario" in archive.read("filtro.csv")


def test_approved_source_manifest_matches_all_tracked_files():
    manifest = payroll.verificar_manifiesto_fuentes()
    assert manifest["integrity_ok"] is True
    assert manifest["application_version"] == payroll.VERSION_APP
    assert len(manifest["sources"]) == len(payroll.FUENTES_CONTROLADAS)
    assert all(len(source["actual_sha256"]) == 64 for source in manifest["sources"])


def test_export_preflight_accepts_consistent_data_and_blocks_invariant_failures(app_namespace):
    proposal, _, records, _ = payroll.extraer_propuesta_plantilla()
    previous = payroll.cargar_salarios_res29()
    scale, detail = app_namespace["crear_tablas"](1.25, proposal, records, previous)
    manifest = {**payroll.verificar_manifiesto_fuentes(), "integrity_ok": True}

    checks, ready = app_namespace["validar_preexportacion"](
        scale, detail, 1.25, manifest, len(records)
    )
    assert ready is True
    assert set(checks["Resultado"]) == {"Correcto"}

    duplicate_ids = detail.copy()
    duplicate_ids.loc[duplicate_ids.index[-1], "Id."] = 1
    failed_ids, ids_ready = app_namespace["validar_preexportacion"](
        scale, duplicate_ids, 1.25, manifest, len(records)
    )
    assert ids_ready is False
    assert failed_ids.set_index("Control").loc["Id. únicos y consecutivos", "Resultado"] == "ERROR"

    _, factor_ready = app_namespace["validar_preexportacion"](
        scale, detail, 2.0, manifest, len(records)
    )
    assert factor_ready is False
    _, factor_range_ready = app_namespace["validar_preexportacion"](
        scale, detail, 6.5, manifest, len(records)
    )
    assert factor_range_ready is False
    _, count_ready = app_namespace["validar_preexportacion"](
        scale, detail, 1.25, manifest, len(records) + 1
    )
    assert count_ready is False

    _, manifest_ready = app_namespace["validar_preexportacion"](
        scale, detail, 1.25, {**manifest, "integrity_ok": False}, len(records)
    )
    assert manifest_ready is False


def test_complete_export_package_contains_outputs_parameters_and_verifiable_hashes(app_namespace, tmp_path):
    proposal, _, records, _, findings = payroll.extraer_propuesta_plantilla_con_calidad()
    detail = app_namespace["crear_tablas"](
        1.25, proposal, records, payroll.cargar_salarios_res29()
    )[1]
    review = app_namespace["crear_tabla_revision_calidad"](findings)
    review.loc[0, "Estado de revisión"] = "En revisión"
    review.loc[0, "Nota de revisión"] = "Validación documentada."
    manifest = {"application_version": "test", "sources": []}
    package = app_namespace["crear_paquete_completo"](
        b"xlsx-output", b"pdf-output", b"approval-draft", 1.25, detail, review, manifest
    )
    package_path = tmp_path / "app_generated_package.zip"
    package_path.write_bytes(package)
    assert verificar_paquete(package_path)["ok"] is True

    with zipfile.ZipFile(BytesIO(package)) as archive:
        names = set(archive.namelist())
        assert {
            "Propuesta_Escala_Resolucion87_1.25x.xlsx",
            "Propuesta_Escala_Resolucion87_1.25x.pdf",
            "Informe_Aprobacion_Capital_Humano_1.25x.pdf",
            "Parametros_escenario.json",
            "Incidencias_calidad_revisadas.csv",
            "manifiesto_fuentes.json",
            "hashes_salidas.json",
        }.issubset(names)
        parameters = json.loads(archive.read("Parametros_escenario.json"))
        assert parameters["factor"] == 1.25
        assert parameters["factor_cap"] == 6
        assert parameters["source_cargo_rows"] == 183
        assert parameters["approval_document_status"] == "DRAFT — NOT APPROVED"
        hashes = json.loads(archive.read("hashes_salidas.json"))
        assert hashes["algorithm"] == "SHA-256"
        assert {entry["file"] for entry in hashes["files"]} == names - {"hashes_salidas.json"}
        for entry in hashes["files"]:
            content = archive.read(entry["file"])
            assert len(content) == entry["bytes"]
            assert hashlib.sha256(content).hexdigest() == entry["sha256"]
        exported_review = pd.read_csv(archive.open("Incidencias_calidad_revisadas.csv"))
        assert exported_review.loc[0, "Estado de revisión"] == "En revisión"
        assert exported_review.loc[0, "Nota de revisión"] == "Validación documentada."


def test_quality_review_rows_have_stable_ids_and_export_manual_annotations(app_namespace):
    findings = [{
        "Tipo de incidencia": "Salario inválido",
        "Hoja de origen": "Hoja de prueba",
        "Fila de origen": 7,
        "Cargo": "Técnico de prueba",
        "Grupo": "XII",
        "Valor detectado": "n/a",
        "Detalle": "El valor de prueba no es numérico.",
    }]
    first = app_namespace["crear_tabla_revision_calidad"](findings)
    second = app_namespace["crear_tabla_revision_calidad"](findings)
    assert first["ID incidencia"].tolist() == second["ID incidencia"].tolist()
    assert first["ID incidencia"].is_unique
    assert set(first["Estado de revisión"]) == {"Pendiente"}
    assert set(first["Nota de revisión"]) == {""}

    first.loc[0, "Estado de revisión"] = "Revisada"
    first.loc[0, "Nota de revisión"] = "Confirmado con el área de origen."
    package = app_namespace["crear_zip_csv_con_manifiesto"](
        first, "Incidencias_calidad_revisadas.csv", payroll.verificar_manifiesto_fuentes()
    )
    with zipfile.ZipFile(BytesIO(package)) as archive:
        exported = pd.read_csv(archive.open("Incidencias_calidad_revisadas.csv"))
    assert exported.loc[0, "Estado de revisión"] == "Revisada"
    assert exported.loc[0, "Nota de revisión"] == "Confirmado con el área de origen."
    assert exported.loc[0, "ID incidencia"] == first.loc[0, "ID incidencia"]


def test_quality_findings_zip_contains_source_locations_and_manifest(app_namespace):
    findings = [{
        "Tipo de incidencia": "Salario inválido",
        "Hoja de origen": "P-2  Aguas Claras",
        "Fila de origen": 44,
        "Cargo": "Cargo de prueba",
        "Grupo": "XII",
        "Valor detectado": "n/a",
        "Detalle": "Hallazgo sintético para validar trazabilidad del ZIP.",
    }]
    package = app_namespace["crear_zip_csv_con_manifiesto"](
        pd.DataFrame(findings), "incidencias.csv", payroll.verificar_manifiesto_fuentes()
    )
    with zipfile.ZipFile(BytesIO(package)) as archive:
        assert set(archive.namelist()) == {"incidencias.csv", "manifiesto_fuentes.json", "LEEME.txt"}
        csv = archive.read("incidencias.csv").decode("utf-8-sig")
        assert "Hoja de origen,Fila de origen" in csv
        assert "P-2  Aguas Claras,44" in csv
        manifest = json.loads(archive.read("manifiesto_fuentes.json"))
        assert manifest["integrity_ok"] is True


def test_checked_in_reference_exports_include_verified_provenance():
    manifest = payroll.verificar_manifiesto_fuentes()
    example_xlsx = APP_DIR / "Escala_Salarial_3210_Nueva.xlsx"
    workbook = load_workbook(example_xlsx, read_only=True, data_only=True)
    assert "Trazabilidad SHA-256" in workbook.sheetnames
    assert workbook["Trazabilidad SHA-256"]["B3"].value == "VERIFICADA"
    detail_sheet = workbook["Detalle Plantilla 87"]
    assert detail_sheet["A1"].value == "Id."
    assert detail_sheet["B1"].value == "Hoja fuente"
    assert detail_sheet["A2"].value == 1
    assert detail_sheet["A184"].value == 183

    example_pdf = APP_DIR / "Escala_Salarial_3210_Nueva.pdf"
    pdf_text = "\n".join(page.extract_text() or "" for page in PdfReader(str(example_pdf)).pages)
    assert "TRAZABILIDAD DE FUENTES" in pdf_text
    assert manifest["sources"][0]["actual_sha256"] in pdf_text


def test_modified_expected_hash_is_detected(tmp_path):
    approved = payroll.crear_manifiesto_fuentes()
    approved["sources"][0]["sha256"] = "0" * 64
    test_manifest = tmp_path / "source_manifest.json"
    test_manifest.write_text(json.dumps(approved), encoding="utf-8")

    result = payroll.verificar_manifiesto_fuentes(test_manifest)
    assert result["integrity_ok"] is False
    assert result["sources"][0]["status"] == "MODIFICADA"


def test_active_scenario_excel_and_pdf_include_traceability(app_namespace):
    proposal, _, records, _ = payroll.extraer_propuesta_plantilla()
    previous = payroll.cargar_salarios_res29()
    scale, details = app_namespace["crear_tablas"](1.25, proposal, records, previous)
    manifest = payroll.verificar_manifiesto_fuentes()
    workbook_bytes = app_namespace["crear_excel"](scale, details, 1.25, manifest)
    workbook = load_workbook(BytesIO(workbook_bytes), data_only=True)
    assert "Trazabilidad SHA-256" in workbook.sheetnames
    assert workbook["Trazabilidad SHA-256"]["B3"].value == "VERIFICADA"
    detail_sheet = workbook["Plantilla Resolución 87"]
    assert [detail_sheet.cell(3, column).value for column in range(1, 4)] == [
        "Id.", "Hoja de origen", "Fila de origen",
    ]
    assert detail_sheet["A4"].value == 1
    assert detail_sheet[f"A{len(details) + 3}"].value == len(details)
    assert detail_sheet.max_column == 9
    assert detail_sheet["G4"].number_format == '0.00"×"'

    pdf_bytes = app_namespace["crear_pdf"](scale, details, 1.25, manifest)
    assert pdf_bytes.startswith(b"%PDF-")
    pdf_pages = PdfReader(BytesIO(pdf_bytes)).pages
    page_texts = [page.extract_text() or "" for page in pdf_pages]
    pdf_text = "\n".join(page_texts)
    assert len(pdf_pages) == 10
    assert all(text.strip() for text in page_texts)
    detail_page = next(text for text in page_texts if "PLANTILLA PROPUESTA · DETALLE DE CARGOS" in text)
    detail_lines = [line for line in detail_page.splitlines() if line.strip()]
    assert detail_lines[0] == "PLANTILLA PROPUESTA · DETALLE DE CARGOS"
    assert detail_lines[1].startswith("Fuente: propuesta de plantilla")
    assert page_texts[-1].startswith("TRAZABILIDAD DE FUENTES")
    assert all(f"Página {number} de {len(pdf_pages)}" in text for number, text in enumerate(page_texts, start=1))
    assert "Id." in pdf_text and "Hoja" in pdf_text
    assert manifest["sources"][0]["actual_sha256"] in pdf_text


def _xlsx_template_bytes(rows):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Plantilla nueva"
    worksheet.append(["Cargo", "G.E.A", "Salario Propuesto"])
    for row in rows:
        worksheet.append(row)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_uploaded_xlsx_is_only_previewed_until_explicit_confirmation():
    payload = _xlsx_template_bytes([
        ["Analista principal", "I", 3210],
        ["Especialista técnico", "II", 3365],
    ])
    result = payroll.validar_plantilla_subida("nueva.xlsx", payload)

    assert result["ready"] is True
    assert result["record_count"] == 2
    assert result["preview"] == result["records"]
    assert result["proposal"] == {"I": 3210.0, "II": 3365.0}
    assert result["size_bytes"] == len(payload)
    assert len(result["sha256"]) == 64
    assert result["errors"] == []


def test_uploaded_template_preserves_possible_duplicates_and_blocks_bad_salaries():
    duplicate_payload = _xlsx_template_bytes([
        ["Especialista principal", "I", 3210],
        ["Especialista principal", "I", 3210],
    ])
    duplicate_result = payroll.validar_plantilla_subida("repetida.xlsx", duplicate_payload)
    assert duplicate_result["ready"] is True
    assert len(duplicate_result["records"]) == 2
    assert any("posibles duplicados" in warning for warning in duplicate_result["warnings"])
    assert any(finding["Tipo de incidencia"] == "Posible duplicado" for finding in duplicate_result["findings"])

    invalid_payload = _xlsx_template_bytes([
        ["Salario cero", "I", 0],
        ["Salario vacío", "II", None],
    ])
    invalid_result = payroll.validar_plantilla_subida("invalida.xlsx", invalid_payload)
    assert invalid_result["ready"] is False
    assert invalid_result["record_count"] == 0
    assert sum(finding["Tipo de incidencia"] in {"Salario ausente", "Salario inválido"} for finding in invalid_result["findings"]) == 2
    assert any("salario ausente o inválido" in error for error in invalid_result["errors"])


def test_uploaded_template_rejects_books_without_expected_group_columns():
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["Cargo", "Grupo", "Importe"])
    worksheet.append(["Analista", "I", 3210])
    output = BytesIO()
    workbook.save(output)

    result = payroll.validar_plantilla_subida("sin_columnas.xlsx", output.getvalue())
    assert result["ready"] is False
    assert result["template_sheets"] == 0
    assert any("encabezados G.E." in error for error in result["errors"])


def test_uploaded_template_rejects_unsupported_malformed_and_oversized_files():
    assert not payroll.validar_plantilla_subida("plantilla.csv", b"data")["ready"]
    malformed = payroll.validar_plantilla_subida("plantilla.xlsx", b"not an xlsx")
    assert not malformed["ready"]
    assert malformed["errors"]
    oversized = payroll.validar_plantilla_subida(
        "plantilla.xls", b"\x00" * (payroll.MAX_PLANTILLA_UPLOAD_BYTES + 1)
    )
    assert not oversized["ready"]
    assert any("15 MiB" in error for error in oversized["errors"])


def test_uploaded_xlsx_rejects_extreme_compression_ratio():
    payload = BytesIO()
    with zipfile.ZipFile(payload, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("xl/worksheets/sheet1.xml", "0" * 40_000)
    result = payroll.validar_plantilla_subida("comprimida.xlsx", payload.getvalue())
    assert result["ready"] is False
    assert any("compresión permitido" in error for error in result["errors"])


def test_uploaded_xls_preview_reads_all_source_rows_without_mutating_the_workbook():
    source = payroll.RUTA_PLANTILLA.read_bytes()
    before = payroll._sha256(payroll.RUTA_PLANTILLA)
    result = payroll.validar_plantilla_subida(payroll.RUTA_PLANTILLA.name, source)

    assert result["ready"] is True
    assert result["record_count"] == 183
    assert result["sha256"] == before
    assert payroll._sha256(payroll.RUTA_PLANTILLA) == before


def test_missing_values_are_blank_in_detail_excel_and_pdf(app_namespace):
    proposal, _, records, _ = payroll.extraer_propuesta_plantilla()
    previous = payroll.cargar_salarios_res29()
    scale, details = app_namespace["crear_tablas"](1.25, proposal, records, previous)
    details.loc[details.index[0], "Hoja de origen"] = "None"
    details.loc[details.index[0], "Cargo"] = None
    manifest = payroll.verificar_manifiesto_fuentes()
    incomplete = {
        **manifest,
        "application_version": None,
        "checked_at_utc": None,
        "integrity_ok": False,
        "sources": [dict(source) for source in manifest["sources"]],
    }
    incomplete["sources"][0].update({
        "label": None, "status": None, "path": None,
        "expected_sha256": None, "actual_sha256": None, "size_bytes": None,
    })

    xlsx = app_namespace["crear_excel"](scale, details, 1.25, incomplete)
    detail_sheet = load_workbook(BytesIO(xlsx), data_only=True)["Plantilla Resolución 87"]
    assert detail_sheet["B4"].value is None
    assert detail_sheet["D4"].value is None

    pdf = app_namespace["crear_pdf"](scale, details, 1.25, incomplete)
    text = "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(pdf)).pages)
    assert "None" not in text
    assert "REVISIÓN REQUERIDA" in text


def test_approval_report_is_a_draft_with_methodology_hashes_and_signature_fields(app_namespace):
    proposal, _, records, anomalies = payroll.extraer_propuesta_plantilla()
    previous = payroll.cargar_salarios_res29()
    scale, details = app_namespace["crear_tablas"](1.25, proposal, records, previous)
    manifest = payroll.verificar_manifiesto_fuentes()
    pdf_bytes = app_namespace["crear_informe_aprobacion_pdf"](
        scale, details, 1.25, manifest, anomalies
    )
    reader = PdfReader(BytesIO(pdf_bytes))
    report_text = "\n".join(page.extract_text() or "" for page in reader.pages)

    assert len(reader.pages) == 3
    assert "INFORME DE METODOLOGÍA Y APROBACIÓN" in report_text
    assert "BORRADOR" in report_text and "PENDIENTE DE REVISIÓN" in report_text
    assert "Resolución 14/2026" in report_text and "Resolución 15/2026" in report_text
    assert "XXXII" in report_text and "6×" in report_text
    assert "No se calcula el impacto presupuestario total" in report_text
    assert "SHA-256 aprobado" in report_text
    for source in manifest["sources"]:
        assert source["expected_sha256"] in report_text
        assert source["actual_sha256"] in report_text
    assert "PREPARÓ" in report_text and "REVISÓ" in report_text and "AUTORIZÓ" in report_text
    assert "Firma manuscrita" in report_text
    assert "Aprobar propuesta" in report_text

    header_positions = []

    def collect_group_table_headers(text, cm, tm, fontdict, fontsize):
        label = (text or "").strip()
        if label in {"Grupo", "Base/ref.", "Factor", "Propuesta"}:
            header_positions.append((label, tm[4]))

    reader.pages[0].extract_text(visitor_text=collect_group_table_headers)
    assert [label for label, _ in header_positions] == [
        "Grupo", "Base/ref.", "Factor", "Propuesta",
        "Grupo", "Base/ref.", "Factor", "Propuesta",
    ]
    page_width = float(reader.pages[0].mediabox.width)
    assert all(0 <= x < page_width for _, x in header_positions)
    assert header_positions[0][1] < page_width / 4

    unverified = {
        **manifest,
        "integrity_ok": False,
        "sources": [dict(source) for source in manifest["sources"]],
    }
    unverified["sources"][0]["status"] = "MODIFICADA"
    unverified_pdf = app_namespace["crear_informe_aprobacion_pdf"](
        scale, details, 1.25, unverified, anomalies
    )
    unverified_text = "\n".join(
        page.extract_text() or "" for page in PdfReader(BytesIO(unverified_pdf)).pages
    )
    assert "REVISIÓN DE FUENTES REQUERIDA" in unverified_text
    assert "MODIFICADA" in unverified_text
