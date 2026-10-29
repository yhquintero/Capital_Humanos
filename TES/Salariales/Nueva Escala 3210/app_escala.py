# -*- coding: utf-8 -*-
"""Interfaz interactiva para recalcular e imprimir la propuesta salarial."""
from __future__ import annotations

import hashlib
import html
import importlib.util
import io
import json
import math
import os
import platform
import sys
import zipfile
from datetime import date, datetime, timezone
from pathlib import Path

import altair as alt
import pandas as pd
import streamlit as st
from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

CARPETA = Path(__file__).resolve().parent
RAIZ_REPO = CARPETA.parents[2]
if str(CARPETA) not in sys.path:
    sys.path.insert(0, str(CARPETA))
import i18n as ui

RUTA_GENERADOR = CARPETA / "generar_escala.py"
_SPEC_GENERADOR = importlib.util.spec_from_file_location(
    "_capital_humano_generar_escala", RUTA_GENERADOR
)
if _SPEC_GENERADOR is None or _SPEC_GENERADOR.loader is None:
    raise ImportError(f"No se pudo cargar el motor salarial desde {RUTA_GENERADOR}")
base = importlib.util.module_from_spec(_SPEC_GENERADOR)
sys.modules[_SPEC_GENERADOR.name] = base
_SPEC_GENERADOR.loader.exec_module(base)
if not hasattr(base, "FACTOR_MAXIMO_ABSOLUTO"):
    raise ImportError(
        "El motor salarial junto a la aplicación no define FACTOR_MAXIMO_ABSOLUTO: "
        f"{RUTA_GENERADOR}. Verifica que ambos archivos sean de la misma versión."
    )

# Permite el comando intuitivo `python app_escala.py` sin ejecutar widgets en
# modo bare. La ejecución normal de Streamlit ya tiene ScriptRunContext y sigue
# directamente al renderizado de la aplicación.
if __name__ == "__main__":
    from streamlit.runtime.scriptrunner import get_script_run_ctx

    if get_script_run_ctx(suppress_warning=True) is None:
        from streamlit.web import cli as stcli

        sys.argv = ["streamlit", "run", str(Path(__file__).resolve()), *sys.argv[1:]]
        raise SystemExit(stcli.main())

AZUL = "15324B"
VERDE = "16866B"
DORADO = "E6B85C"
FUENTE = base.FUENTE
FUENTE_B = base.FUENTE_B

UI_LANGUAGE_LABEL = st.session_state.get("idioma_interfaz", "Español")
UI_LANGUAGE = "en" if UI_LANGUAGE_LABEL == "English" else "es"


def T(text: str, **values) -> str:
    return ui.translate(text, UI_LANGUAGE, **values)


st.set_page_config(
    page_title="Human Capital | Dashboard" if UI_LANGUAGE == "en" else "Capital Humano | Panel de Control",
    page_icon="📊",
    layout="wide",
)


@st.cache_data(show_spinner=T("Leyendo la plantilla y verificando los grupos…"))
def _cargar_datos_oficiales():
    propuesta, contadores, registros, anomalias, hallazgos_calidad = base.extraer_propuesta_plantilla_con_calidad()
    anterior = base.cargar_salarios_res29()
    return propuesta, contadores, registros, anomalias, anterior, hallazgos_calidad


def cargar_datos(plantilla_importada_sha256=None):
    plantilla_temporal = st.session_state.get("plantilla_importada")
    if plantilla_importada_sha256:
        if not plantilla_temporal or plantilla_temporal.get("sha256") != plantilla_importada_sha256:
            raise ValueError("La carga temporal cambió durante la validación de la sesión.")
        # Las cargas salariales temporales nunca pasan por el caché global de Streamlit.
        return (
            plantilla_temporal["proposal"],
            plantilla_temporal["counters"],
            plantilla_temporal["records"],
            plantilla_temporal["anomalies"],
            base.cargar_salarios_res29(),
            plantilla_temporal["findings"],
        )
    return _cargar_datos_oficiales()


def cargar_manifiesto():
    manifiesto = base.manifiesto_para_salida()
    plantilla_temporal = st.session_state.get("plantilla_importada")
    if not plantilla_temporal:
        return manifiesto
    manifiesto = {**manifiesto, "integrity_ok": False, "sources": list(manifiesto.get("sources", []))}
    manifiesto["sources"].append({
        "id": "plantilla_subida_sesion",
        "label": "Plantilla importada (solo sesión; pendiente de aprobación)",
        "path": f"Carga de sesión: {plantilla_temporal['filename']}",
        "status": "NO APROBADA · SOLO SESIÓN",
        "expected_sha256": "",
        "actual_sha256": plantilla_temporal["sha256"],
        "size_bytes": plantilla_temporal["size_bytes"],
    })
    return manifiesto


def descargas_deshabilitadas():
    """Bloquea exportaciones salariales cuando se usa una plantilla no aprobada."""
    return bool(st.session_state.get("plantilla_importada"))


def crear_tablas(factor, propuesta, registros, anterior):
    filas = base.construir_filas(anterior, propuesta)
    escala = []
    for f in filas:
        salario_base = f["plantilla"]
        if salario_base is not None:
            techo = salario_base * base.FACTOR_MAXIMO_ABSOLUTO
            salario = base.aplicar_factor_con_tope(salario_base, factor)
            fuente = "Plantilla Resolución 87"
        else:
            # Los grupos sin cargos propuestos conservan la escala oficial vigente;
            # el mínimo del grupo I queda anclado en $ 3,210.
            salario_base = f["oficial"]
            techo = salario_base * base.FACTOR_MAXIMO_ABSOLUTO
            salario = salario_base
            fuente = "Escala oficial Res. 15/2026"
        escala.append({
            "Grupo": f["grupo"],
            "Salario anterior Res. 29/2020 (CUP)": f["anterior"],
            "Escala oficial Res. 15/2026 (CUP)": f["oficial"],
            "Base propuesta / referencia (CUP)": salario_base,
            "Factor aplicado": factor if f["plantilla"] is not None else 1.0,
            "Tope máximo (6× base) (CUP)": techo,
            "Nueva propuesta (CUP)": salario,
            "Fuente del salario": fuente,
        })
    detalle = []
    for item_id, reg in enumerate(registros, start=1):
        base_salario = float(reg["salario"])
        tope = base_salario * base.FACTOR_MAXIMO_ABSOLUTO
        ajustado = base.aplicar_factor_con_tope(base_salario, factor)
        # Invariante de negocio: nunca superar seis veces el valor de origen.
        assert ajustado <= tope + 1e-9
        detalle.append({
            "Id.": item_id,
            "Hoja de origen": _texto_visible(reg.get("hoja")),
            "Fila de origen": reg.get("fila"),
            "Cargo": _texto_visible(reg.get("cargo")),
            "Grupo propuesto": _texto_visible(reg.get("grupo")),
            "Salario original (CUP)": base_salario,
            "Factor aplicado": factor,
            "Tope máximo 6× (CUP)": tope,
            "Salario ajustado (CUP)": ajustado,
        })
    return pd.DataFrame(escala), pd.DataFrame(detalle)


def crear_comparacion_escenarios(factor_a, factor_b, propuesta, anterior):
    """Devuelve una tabla auditable de diferencias salariales por grupo."""
    return pd.DataFrame(base.comparar_escenarios(factor_a, factor_b, propuesta, anterior))


def crear_comparacion_cargos(factor_a, factor_b, registros):
    """Compara dos factores por cargo, preservando Id. y coordenadas de origen."""
    factores = (float(factor_a), float(factor_b))
    if any(not math.isfinite(value) or not 1 <= value <= base.FACTOR_MAXIMO_ABSOLUTO for value in factores):
        raise ValueError("Los factores deben estar entre 1× y 6×.")
    columnas = [
        "Id.", "Hoja de origen", "Fila de origen", "Cargo", "Grupo propuesto",
        "Salario fuente (CUP)", "Factor A", "Salario A (CUP)", "Factor B",
        "Salario B (CUP)", "Variación B − A (CUP)", "Variación relativa",
    ]
    filas = []
    for item_id, registro in enumerate(registros, start=1):
        salario = float(registro["salario"])
        salario_a = base.aplicar_factor_con_tope(salario, factores[0])
        salario_b = base.aplicar_factor_con_tope(salario, factores[1])
        variacion = salario_b - salario_a
        filas.append({
            "Id.": item_id,
            "Hoja de origen": _texto_visible(registro.get("hoja")),
            "Fila de origen": registro.get("fila"),
            "Cargo": _texto_visible(registro.get("cargo")),
            "Grupo propuesto": _texto_visible(registro.get("grupo")),
            "Salario fuente (CUP)": salario,
            "Factor A": factores[0],
            "Salario A (CUP)": salario_a,
            "Factor B": factores[1],
            "Salario B (CUP)": salario_b,
            "Variación B − A (CUP)": variacion,
            "Variación relativa": variacion / salario_a if salario_a else 0.0,
        })
    return pd.DataFrame(filas, columns=columnas)


def crear_excel_comparacion_cargos(comparacion: pd.DataFrame, factor_a: float, factor_b: float, manifiesto: dict) -> bytes:
    """Crea un libro filtrable de comparación A/B por cargo y con trazabilidad."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Comparación por cargo"
    ws.sheet_view.showGridLines = False
    ultima_columna = get_column_letter(len(comparacion.columns))
    ws.append(["COMPARACIÓN DE ESCENARIOS · DETALLE POR CARGO"])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(comparacion.columns))
    ws.append([f"Escenario A: {factor_a:.2f}× · Escenario B: {factor_b:.2f}× · CUP · {date.today():%d/%m/%Y}"])
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(comparacion.columns))
    ws.append(list(comparacion.columns))
    for values in comparacion.itertuples(index=False, name=None):
        ws.append([None if _es_vacio(value) else value for value in values])

    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=AZUL)
        cell.font = Font(name="Calibri", size=15, bold=True, color="FFFFFF")
    for cell in ws[2]:
        cell.fill = PatternFill("solid", fgColor="EAF0F4")
        cell.font = Font(name="Calibri", size=10, italic=True, color=AZUL)
    for cell in ws[3]:
        cell.fill = PatternFill("solid", fgColor=VERDE)
        cell.font = Font(name="Calibri", size=9, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 24
    ws.row_dimensions[3].height = 36
    borde = Side(style="hair", color="CBD5DC")
    for row in ws.iter_rows(min_row=4, max_row=ws.max_row):
        for cell in row:
            cell.border = Border(bottom=borde)
            cell.alignment = Alignment(vertical="center", wrap_text=cell.column in (2, 4))
        for index, nombre in enumerate(comparacion.columns, start=1):
            if "CUP" in nombre:
                row[index - 1].number_format = '"$ "#,##0.00;[Red]"$ "-#,##0.00'
            elif nombre.startswith("Factor "):
                row[index - 1].number_format = '0.00"×"'
            elif nombre == "Variación relativa":
                row[index - 1].number_format = "0.00%"
            elif nombre in {"Id.", "Fila de origen"}:
                row[index - 1].number_format = "0"
    anchos = {
        "Id.": 9, "Hoja de origen": 26, "Fila de origen": 12, "Cargo": 50,
        "Grupo propuesto": 17, "Salario fuente (CUP)": 22, "Factor A": 13,
        "Salario A (CUP)": 20, "Factor B": 13, "Salario B (CUP)": 20,
        "Variación B − A (CUP)": 25, "Variación relativa": 20,
    }
    for index, nombre in enumerate(comparacion.columns, start=1):
        ws.column_dimensions[get_column_letter(index)].width = anchos.get(nombre, 18)
    if ws.max_row > 3:
        tabla = Table(displayName="ComparacionPorCargo", ref=f"A3:{ultima_columna}{ws.max_row}")
        tabla.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(tabla)
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:{ultima_columna}{ws.max_row}"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:3"
    ws.print_area = f"A1:{ultima_columna}{ws.max_row}"
    agregar_hoja_trazabilidad(wb, manifiesto)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def crear_excel_comparacion(comparacion: pd.DataFrame, factor_a: float, factor_b: float, manifiesto: dict) -> bytes:
    """Crea un Excel de comparación, listo para filtrar e imprimir en A4."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Comparación por grupo"
    ws.sheet_view.showGridLines = False
    ws.append(["COMPARACIÓN DE ESCENARIOS · CAPITAL HUMANO"])
    ws.merge_cells("A1:F1")
    ws.append([f"Escenario A: {factor_a:.2f}× · Escenario B: {factor_b:.2f}× · CUP · {date.today():%d/%m/%Y}"])
    ws.merge_cells("A2:F2")
    ws.append(list(comparacion.columns))
    for fila in comparacion.itertuples(index=False, name=None):
        ws.append([None if _es_vacio(value) else value for value in fila])
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=AZUL)
        cell.font = Font(name="Calibri", size=15, bold=True, color="FFFFFF")
    for cell in ws[2]:
        cell.fill = PatternFill("solid", fgColor="EAF0F4")
        cell.font = Font(name="Calibri", size=10, italic=True, color=AZUL)
    for cell in ws[3]:
        cell.fill = PatternFill("solid", fgColor=VERDE)
        cell.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[3].height = 34
    edge = Side(style="hair", color="CBD5DC")
    for row in ws.iter_rows(min_row=4, max_row=ws.max_row):
        for cell in row:
            cell.border = Border(bottom=edge)
            cell.alignment = Alignment(vertical="center", wrap_text=cell.column in (1, 6))
        for index in (2, 3, 4):
            row[index - 1].number_format = '"$ "#,##0.00'
        row[4].number_format = "0.00%"
    for column, width in enumerate((13, 22, 22, 24, 22, 45), start=1):
        ws.column_dimensions[chr(64 + column)].width = width
    if ws.max_row > 3:
        table = Table(displayName="ComparacionEscenarios", ref=f"A3:F{ws.max_row}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(table)
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:F{ws.max_row}"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.print_title_rows = "1:3"
    ws.print_area = f"A1:F{ws.max_row}"
    agregar_hoja_trazabilidad(wb, manifiesto)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def validar_preexportacion(escala: pd.DataFrame, detalle: pd.DataFrame, factor: float, manifiesto: dict, cantidad_esperada: int):
    """Valida invariantes de exportación sin tocar ni reparar los datos."""
    controles = []

    def registrar(nombre, correcto, detalle_control):
        controles.append({
            "Control": nombre,
            "Resultado": "Correcto" if correcto else "ERROR",
            "Detalle": detalle_control,
        })

    try:
        factor_numerico = float(factor)
        factor_valido = math.isfinite(factor_numerico) and 1.0 <= factor_numerico <= base.FACTOR_MAXIMO_ABSOLUTO
    except (TypeError, ValueError, OverflowError):
        factor_numerico = math.nan
        factor_valido = False
    registrar("Factor dentro del rango 1×–6×", factor_valido, f"Factor activo: {factor}.")
    registrar(
        "Integridad SHA-256 de las fuentes",
        manifiesto.get("integrity_ok") is True,
        "Todas las fuentes y códigos coinciden con el manifiesto aprobado." if manifiesto.get("integrity_ok") is True else "Hay fuentes/código por revisar; las descargas quedan bloqueadas.",
    )

    escala_correcta = (
        "Grupo" in escala.columns
        and len(escala) == 32
        and escala["Grupo"].nunique(dropna=False) == 32
    )
    registrar("Escala completa", escala_correcta, f"Se esperan 32 grupos; encontrados: {len(escala)}.")
    registrar(
        "Cantidad de filas de detalle",
        len(detalle) == cantidad_esperada,
        f"Registros extraídos: {cantidad_esperada}; filas a exportar: {len(detalle)}.",
    )

    id_correcto = False
    factor_coincide = False
    limite_correcto = False
    campos_completos = False
    columnas_requeridas = {
        "Id.", "Hoja de origen", "Fila de origen", "Grupo propuesto",
        "Salario original (CUP)", "Factor aplicado",
        "Tope máximo 6× (CUP)", "Salario ajustado (CUP)",
    }
    if columnas_requeridas.issubset(detalle.columns):
        id_correcto = detalle["Id."].tolist() == list(range(1, len(detalle) + 1))
        factores = pd.to_numeric(detalle["Factor aplicado"], errors="coerce")
        factor_coincide = math.isfinite(factor_numerico) and factores.notna().all() and all(
            math.isclose(float(value), factor_numerico, rel_tol=0, abs_tol=1e-9)
            for value in factores
        )
        salarios = pd.to_numeric(detalle["Salario ajustado (CUP)"], errors="coerce")
        topes = pd.to_numeric(detalle["Tope máximo 6× (CUP)"], errors="coerce").round(2)
        limite_correcto = salarios.notna().all() and topes.notna().all() and salarios.le(topes + 1e-9).all()
        campos_a_revisar = [
            "Id.", "Hoja de origen", "Fila de origen", "Grupo propuesto",
            "Salario original (CUP)", "Factor aplicado",
            "Tope máximo 6× (CUP)", "Salario ajustado (CUP)",
        ]
        campos_completos = not detalle[campos_a_revisar].isna().any().any()
    registrar("Id. únicos y consecutivos", id_correcto, "Los identificadores deben ser 1…N, sin huecos ni duplicados.")
    registrar("Factor de cada fila coincide", factor_coincide, "Todas las filas deben usar el factor mostrado en el escenario activo.")
    registrar("Tope 6× por cargo", limite_correcto, "Cada salario exportado debe estar bajo su tope individual.")
    registrar("Campos requeridos completos", campos_completos, "Se requieren origen, fila, grupo, base, factor y valores calculados.")
    if not columnas_requeridas.issubset(detalle.columns):
        registrar("Estructura de detalle", False, "Faltan una o más columnas necesarias para generar informes auditables.")

    informe = pd.DataFrame(controles)
    listo = informe["Resultado"].eq("Correcto").all()
    return informe, bool(listo)


def agregar_hoja_trazabilidad(wb: Workbook, manifiesto: dict) -> None:
    """Añade la hoja de trazabilidad compartida con el generador de cálculos."""
    base.agregar_hoja_trazabilidad(wb, manifiesto)


def crear_zip_csv_con_manifiesto(df: pd.DataFrame, csv_name: str, manifiesto: dict) -> bytes:
    """Entrega el CSV intacto junto a su manifiesto legible en un ZIP."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, mode="w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(csv_name, df.to_csv(index=False).encode("utf-8-sig"))
        archive.writestr(
            "manifiesto_fuentes.json",
            json.dumps(manifiesto, ensure_ascii=False, indent=2).encode("utf-8"),
        )
        archive.writestr(
            "LEEME.txt",
            "El CSV conserva sus filas y columnas. El manifiesto JSON identifica las fuentes y compara los hashes observados con los aprobados en el repositorio. SHA-256 detecta diferencias de bytes; no es firma digital ni certificación normativa.\n",
        )
    return buffer.getvalue()


def crear_paquete_completo(excel_bytes, pdf_bytes, aprobacion_bytes, factor, detalle, revision_calidad, manifiesto):
    """Empaqueta los documentos del mismo escenario y un inventario SHA-256."""
    sufijo = f"{float(factor):.2f}x"
    fecha_creacion = datetime.now(timezone.utc).isoformat(timespec="seconds")
    archivos = {
        f"Propuesta_Escala_Resolucion87_{sufijo}.xlsx": excel_bytes,
        f"Propuesta_Escala_Resolucion87_{sufijo}.pdf": pdf_bytes,
        f"Informe_Aprobacion_Capital_Humano_{sufijo}.pdf": aprobacion_bytes,
    }
    parametros = {
        "application_version": manifiesto.get("application_version", base.VERSION_APP),
        "generated_at_utc": fecha_creacion,
        "scenario_date": date.today().isoformat(),
        "factor": float(factor),
        "factor_cap": base.FACTOR_MAXIMO_ABSOLUTO,
        "currency": "CUP",
        "source_cargo_rows": int(len(detalle)),
        "scale_groups": 32,
        "approval_document_status": "DRAFT — NOT APPROVED",
        "approval_document_filename": f"Informe_Aprobacion_Capital_Humano_{sufijo}.pdf",
    }
    archivos["Parametros_escenario.json"] = json.dumps(
        parametros, ensure_ascii=False, indent=2, sort_keys=True
    ).encode("utf-8")
    archivos["Incidencias_calidad_revisadas.csv"] = revision_calidad.to_csv(
        index=False, lineterminator="\n"
    ).encode("utf-8-sig")
    archivos["manifiesto_fuentes.json"] = json.dumps(
        manifiesto, ensure_ascii=False, indent=2, sort_keys=True, default=str
    ).encode("utf-8")
    readme = (
        "PAQUETE DE EXPORTACIÓN · CAPITAL HUMANO\n"
        f"Escenario: {float(factor):.2f}× · Moneda: CUP · Registros: {len(detalle)}\n"
        "El informe de aprobación es un borrador y no constituye una autorización.\n"
        "Incidencias y notas: revisión manual de esta sesión; no altera la plantilla fuente.\n"
        "Los hashes SHA-256 de los demás archivos se encuentran en hashes_salidas.json.\n"
        "Trate todos los documentos como información salarial sensible.\n"
    ).encode("utf-8")
    archivos["LEEME.txt"] = readme
    inventario_hashes = {
        "algorithm": "SHA-256",
        "generated_at_utc": fecha_creacion,
        "coverage": "All package files except this hashes_salidas.json file.",
        "files": [
            {"file": nombre, "bytes": len(contenido), "sha256": hashlib.sha256(contenido).hexdigest()}
            for nombre, contenido in sorted(archivos.items())
        ],
    }
    archivos["hashes_salidas.json"] = json.dumps(
        inventario_hashes, ensure_ascii=False, indent=2, sort_keys=True
    ).encode("utf-8")
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for nombre, contenido in archivos.items():
            archive.writestr(nombre, contenido)
    return buffer.getvalue()


def crear_tabla_revision_calidad(hallazgos):
    """Prepara incidencias con un ID estable y campos de revisión humana editables."""
    columnas_reporte = [
        "Tipo de incidencia", "Hoja de origen", "Fila de origen", "Cargo",
        "Grupo", "Valor detectado", "Detalle",
    ]
    reporte = pd.DataFrame(hallazgos, columns=columnas_reporte)
    reporte.insert(0, "ID incidencia", [
        "Q-" + hashlib.sha256(
            "|".join(str(fila.get(columna, "")) for columna in columnas_reporte).encode("utf-8")
        ).hexdigest()[:12].upper()
        for fila in hallazgos
    ])
    reporte["Estado de revisión"] = "Pendiente"
    reporte["Nota de revisión"] = ""
    return reporte


def _es_vacio(value) -> bool:
    """Detecta nulos y marcadores textuales de nulo para no imprimir `None`."""
    if value is None:
        return True
    if isinstance(value, str) and value.strip().casefold() in {"none", "nan", "<na>"}:
        return True
    try:
        return bool(pd.isna(value))
    except (TypeError, ValueError):
        return False


def _texto_visible(value, por_defecto: str = "") -> str:
    """Convierte datos ausentes en texto vacío o una etiqueta explícita."""
    return por_defecto if _es_vacio(value) else str(value)


def crear_excel(escala: pd.DataFrame, detalle: pd.DataFrame, factor: float, manifiesto: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Escala interactiva"
    ws.sheet_view.showGridLines = False
    ws.append(["PROPUESTA SALARIAL · CAPITAL HUMANO"])
    ws.append([f"Factor: {factor:.2f}× · Límite: 6× del salario fuente · Generado: {date.today():%d/%m/%Y}"])
    ws.append(list(escala.columns))
    for row in escala.itertuples(index=False, name=None):
        ws.append([None if _es_vacio(value) else value for value in row])
    ultima_columna_escala = get_column_letter(len(escala.columns))
    _estilizar_hoja(
        ws, 3, len(escala) + 3, "TablaEscala",
        [12, 25, 26, 27, 16, 25, 24, 32],
        money_cols={2, 3, 4, 6, 7}, factor_cols={5}, wrap_cols={8},
    )
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:{ultima_columna_escala}{len(escala) + 3}"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.print_title_rows = "1:3"
    ws.print_area = f"A1:{ultima_columna_escala}{len(escala)+3}"

    wd = wb.create_sheet("Plantilla Resolución 87")
    wd.sheet_view.showGridLines = False
    wd.append(["PROPUESTA ACTUALIZADA DE PLANTILLA · RESOLUCIÓN 87"])
    wd.append([f"Factor: {factor:.2f}× · Máximo permitido: 6× del salario original · Generado: {date.today():%d/%m/%Y}"])
    wd.append(list(detalle.columns))
    for row in detalle.itertuples(index=False, name=None):
        wd.append([None if _es_vacio(value) else value for value in row])
    ultima_columna_detalle = get_column_letter(len(detalle.columns))
    _estilizar_hoja(
        wd, 3, len(detalle) + 3, "DetallePlantilla",
        [9, 28, 11, 54, 18, 24, 13, 23, 25],
        money_cols={6, 8, 9}, factor_cols={7}, wrap_cols={4},
    )
    wd.freeze_panes = "A4"
    wd.auto_filter.ref = f"A3:{ultima_columna_detalle}{len(detalle) + 3}"
    wd.sheet_properties.pageSetUpPr.fitToPage = True
    wd.page_setup.orientation = "landscape"
    wd.page_setup.paperSize = wd.PAPERSIZE_A4
    wd.page_setup.fitToWidth = 1
    wd.page_setup.fitToHeight = 0
    wd.print_title_rows = "1:3"
    wd.print_area = f"A1:{ultima_columna_detalle}{len(detalle)+3}"
    agregar_hoja_trazabilidad(wb, manifiesto)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _estilizar_hoja(
    ws,
    header_row,
    last_row,
    table_name,
    widths,
    *,
    money_cols=(),
    factor_cols=(),
    wrap_cols=(),
):
    ultima_columna = get_column_letter(len(widths))
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(widths))
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(widths))
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=AZUL)
        cell.font = Font(name="Calibri", size=15, bold=True, color="FFFFFF")
        cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 30
    for cell in ws[2]:
        cell.fill = PatternFill("solid", fgColor="EAF0F4")
        cell.font = Font(name="Calibri", size=10, italic=True, color=AZUL)
    for cell in ws[header_row]:
        cell.fill = PatternFill("solid", fgColor=VERDE)
        cell.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[header_row].height = 36
    edge = Side(style="hair", color="CBD5DC")
    for row in ws.iter_rows(min_row=header_row + 1, max_row=last_row):
        for cell in row:
            cell.border = Border(bottom=edge)
            cell.alignment = Alignment(vertical="center", wrap_text=cell.column in wrap_cols)
            if cell.column in money_cols and isinstance(cell.value, (int, float)):
                cell.number_format = '"$ "#,##0.00'
            elif cell.column in factor_cols and isinstance(cell.value, (int, float)):
                cell.number_format = '0.00"×"'
    if last_row > header_row:
        tab = Table(displayName=table_name, ref=f"A{header_row}:{ultima_columna}{last_row}")
        tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True,
                                            showColumnStripes=False)
        ws.add_table(tab)
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width


def crear_pdf(escala: pd.DataFrame, detalle: pd.DataFrame, factor: float, manifiesto: dict) -> bytes:
    version_pie = _texto_visible(manifiesto.get("application_version"), base.VERSION_APP)
    fecha_pie = date.today().strftime("%d/%m/%Y")

    class PDFConPie(FPDF):
        def footer(self):
            self.set_y(-9)
            self.set_font("DejaVu", "", 7)
            self.set_text_color(100, 110, 118)
            self.cell(
                0, 5,
                f"Capital Humano · v{version_pie} · {fecha_pie} · Página {self.page_no()} de {{nb}}",
                align="C",
            )

    pdf = PDFConPie(orientation="L", unit="mm", format="A4")
    pdf.add_font("DejaVu", "", FUENTE)
    pdf.add_font("DejaVu", "B", FUENTE_B)
    pdf.set_auto_page_break(auto=True, margin=14)
    pdf.set_margins(10, 10, 10)
    pdf.title = "Propuesta de escala y plantilla salarial"
    pdf.set_author("Capital Humano")

    def cabecera(titulo, subtitulo):
        pdf.set_font("DejaVu", "B", 15)
        pdf.set_text_color(21, 50, 75)
        pdf.cell(0, 8, titulo, new_x="LMARGIN", new_y="NEXT", align="C")
        pdf.set_font("DejaVu", "", 9)
        pdf.set_text_color(70, 85, 95)
        pdf.cell(0, 6, subtitulo, new_x="LMARGIN", new_y="NEXT", align="C")
        pdf.ln(3)

    def tabla(df, titulos, widths, money_cols, factor_cols=(), font_size=7.3):
        pdf.set_font("DejaVu", "B", font_size)
        pdf.set_fill_color(22, 134, 107)
        pdf.set_text_color(255, 255, 255)
        for title, width in zip(titulos, widths):
            pdf.multi_cell(width, 8, str(title), border=1, align="C", fill=True,
                           new_x="RIGHT", new_y="TOP", max_line_height=4)
        pdf.ln(8)
        pdf.set_text_color(25, 39, 49)
        for ridx, row in enumerate(df.itertuples(index=False, name=None)):
            if pdf.get_y() > 185:
                pdf.add_page()
                pdf.set_font("DejaVu", "B", font_size)
                pdf.set_fill_color(22, 134, 107)
                pdf.set_text_color(255, 255, 255)
                for title, width in zip(titulos, widths):
                    pdf.multi_cell(width, 8, str(title), border=1, align="C", fill=True,
                                   new_x="RIGHT", new_y="TOP", max_line_height=4)
                pdf.ln(8)
                pdf.set_text_color(25, 39, 49)
            color_fondo_fila = (244, 247, 248) if ridx % 2 else (255, 255, 255)
            pdf.set_fill_color(*color_fondo_fila)
            pdf.set_font("DejaVu", "", font_size)
            height = 6
            for idx, (value, width) in enumerate(zip(row, widths)):
                if _es_vacio(value):
                    text = ""
                elif isinstance(value, (int, float)) and idx in money_cols:
                    text = f"$ {value:,.2f}"
                elif isinstance(value, (int, float)) and idx in factor_cols:
                    text = f"{value:.2f}×"
                elif isinstance(value, float):
                    text = f"{value:,.2f}"
                else:
                    text = str(value)
                if len(text) > 52 and width < 55:
                    text = text[:49] + "…"
                pdf.cell(width, height, text, border=1, align="R" if idx in money_cols or idx in factor_cols else "L",
                         fill=True, new_x="RIGHT", new_y="TOP")
            pdf.ln(height)

    pdf.add_page()
    cabecera("NUEVA ESCALA SALARIAL", f"Mínimo $ 3,210.00 · Factor de plantilla {factor:.2f}× · Tope absoluto 6× del salario base · CUP")
    cols = ["Grupo", "Anterior\nRes. 29/2020", "Oficial\nRes. 15/2026", "Base", "Factor", "Tope 6×", "Nueva propuesta", "Fuente"]
    tabla(escala, cols, [14, 30, 30, 30, 18, 30, 31, 51], {1, 2, 3, 5, 6}, factor_cols={4})
    pdf.ln(2)
    pdf.set_font("DejaVu", "", 8)
    pdf.multi_cell(0, 4.5, "El grupo I conserva el salario mínimo oficial de $ 3,210.00. Los grupos sin propuesta de cargos conservan la escala oficial. La tabla detallada siguiente aplica el factor a cada salario de la plantilla fuente y limita cada resultado a seis veces ese salario.")
    pdf.add_page()
    cabecera("PLANTILLA PROPUESTA · DETALLE DE CARGOS", f"Fuente: propuesta de plantilla por la resolución 87.xls · {len(detalle)} cargos · Factor {factor:.2f}×")
    det_cols = ["Id.", "Hoja", "Fila", "Cargo", "Grupo", "Original", "Factor", "Tope 6×", "Ajustado"]
    tabla(
        detalle,
        det_cols,
        [10, 28, 12, 72, 15, 27, 17, 27, 28],
        {5, 7, 8},
        factor_cols={6},
        font_size=6.8,
    )
    pdf.add_page()
    estado_general = "INTEGRIDAD VERIFICADA" if manifiesto.get("integrity_ok") else "REVISIÓN REQUERIDA"
    version = _texto_visible(manifiesto.get("application_version"), "desconocida")
    checked_at = _texto_visible(manifiesto.get("checked_at_utc"), "No disponible")
    cabecera("TRAZABILIDAD DE FUENTES", f"Versión {version} · {estado_general}")
    pdf.set_font("DejaVu", "", 8)
    pdf.multi_cell(0, 4.5, f"Verificado en UTC: {checked_at}\nLas huellas SHA-256 comparan los archivos actuales con el manifiesto versionado en docs/source_manifest.json. Coincidir no certifica autoría, firma digital ni vigencia jurídica de la fuente.")
    pdf.ln(2)
    for fuente in manifiesto.get("sources", []):
        estado_fuente = _texto_visible(fuente.get("status"), "NO VERIFICADA")
        etiqueta = _texto_visible(fuente.get("label"), "Fuente sin etiqueta")
        ruta = _texto_visible(fuente.get("path"), "no disponible")
        hash_aprobado = _texto_visible(fuente.get("expected_sha256"), "no disponible")
        hash_actual = _texto_visible(fuente.get("actual_sha256"), "no disponible")
        tamano = _texto_visible(fuente.get("size_bytes"), "no disponible")
        pdf.set_font("DejaVu", "B", 8.5)
        pdf.cell(0, 5, f"{etiqueta} · {estado_fuente}", new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("DejaVu", "", 7)
        pdf.multi_cell(0, 4, f"Ruta: {ruta}\nSHA-256 aprobado: {hash_aprobado}\nSHA-256 actual: {hash_actual} · Tamaño: {tamano} bytes")
        pdf.ln(1)
    if manifiesto.get("error"):
        pdf.set_font("DejaVu", "B", 8)
        pdf.multi_cell(0, 4, f"Error al verificar el manifiesto: {manifiesto['error']}")
    pdf.alias_nb_pages()
    return bytes(pdf.output())


def crear_informe_aprobacion_pdf(
    escala: pd.DataFrame,
    detalle: pd.DataFrame,
    factor: float,
    manifiesto: dict,
    anomalias: list[str] | None = None,
) -> bytes:
    """Genera un informe de revisión imprimible, sin simular una aprobación."""
    pdf = FPDF(orientation="P", unit="mm", format="A4")
    pdf.add_font("DejaVu", "", FUENTE)
    pdf.add_font("DejaVu", "B", FUENTE_B)
    pdf.set_auto_page_break(auto=True, margin=13)
    pdf.set_margins(10, 10, 10)
    pdf.title = "Informe de metodología y solicitud de aprobación"
    pdf.subject = "Documento de trabajo pendiente de revisión y autorización humana"
    pdf.set_author("Capital Humano")

    version = _texto_visible(manifiesto.get("application_version"), base.VERSION_APP)
    integrity_ok = bool(manifiesto.get("integrity_ok"))
    integrity_label = "FUENTES VERIFICADAS" if integrity_ok else "REVISIÓN DE FUENTES REQUERIDA"

    def heading(text: str) -> None:
        pdf.set_font("DejaVu", "B", 10)
        pdf.set_text_color(21, 50, 75)
        pdf.cell(0, 6, text, new_x="LMARGIN", new_y="NEXT")

    def paragraph(text: str, size: float = 7.2, line_height: float = 4.4) -> None:
        pdf.set_font("DejaVu", "", size)
        pdf.set_text_color(35, 55, 70)
        pdf.multi_cell(0, line_height, text, new_x="LMARGIN", new_y="NEXT")

    # Página 1: alcance, escenario y resumen agregado de los 32 grupos.
    pdf.add_page()
    pdf.set_y(13)
    pdf.set_font("DejaVu", "B", 14)
    pdf.set_text_color(21, 50, 75)
    pdf.cell(0, 8, "INFORME DE METODOLOGÍA Y APROBACIÓN", align="C", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("DejaVu", "B", 8)
    pdf.set_text_color(130, 77, 12)
    pdf.set_fill_color(255, 243, 220)
    pdf.cell(0, 7, "BORRADOR · PENDIENTE DE REVISIÓN Y APROBACIÓN HUMANA", border=1, align="C", fill=True, new_x="LMARGIN", new_y="NEXT")
    pdf.ln(3)
    paragraph(
        f"Fecha de generación: {date.today():%d/%m/%Y}    ·    Versión: {version}    ·    "
        f"Factor seleccionado: {factor:.2f}×    ·    Cargos fuente: {len(detalle):,}    ·    "
        f"Integridad: {integrity_label}",
        size=7.6,
        line_height=5,
    )
    pdf.ln(2)
    heading("Propósito y alcance")
    paragraph(
        "Este documento resume un escenario de modelación para facilitar su revisión. No constituye una "
        "resolución, decisión administrativa, autorización salarial ni estimación del gasto total de nómina. "
        "La aplicación no dispone del número aprobado de plazas ocupadas y vacantes."
    )
    pdf.ln(1)
    heading("Método aplicado")
    paragraph(
        "Para los grupos con cargos propuestos, el cálculo es min(salario de origen × factor, salario de origen × 6). "
        "El factor seleccionado se aplica por cargo y nunca supera el límite absoluto de 6×. El grupo I conserva "
        "el salario mínimo de la Resolución 14/2026 ($ 3,210.00 CUP); los grupos sin propuesta conservan la "
        "escala oficial de la Resolución 15/2026 para 44 horas. Las firmas y la decisión se completan "
        "manualmente en la página 3."
    )
    pdf.ln(2)
    heading("Resumen por grupo · CUP")
    pdf.set_font("DejaVu", "", 6.5)
    pdf.set_text_color(80, 94, 104)
    pdf.multi_cell(0, 3.8, "Base/ref. identifica el salario usado como referencia en cada grupo; la propuesta respeta el tope 6×.")
    records = sorted(
        escala.to_dict("records"),
        key=lambda row: base.ROMANO_A_ENTERO.get(str(row.get("Grupo", "")), 999),
    )
    half = (len(records) + 1) // 2
    widths = (15, 30, 15, 35, 15, 30, 15, 35)
    headers = ("Grupo", "Base/ref.", "Factor", "Propuesta", "Grupo", "Base/ref.", "Factor", "Propuesta")
    pdf.set_font("DejaVu", "B", 6.8)
    pdf.set_fill_color(22, 134, 107)
    pdf.set_text_color(255, 255, 255)
    # multi_cell deja el cursor en el borde derecho; reiniciar la tabla en el margen.
    pdf.set_x(pdf.l_margin)
    for title, width in zip(headers, widths):
        pdf.cell(width, 7, title, border=1, align="C", fill=True, new_x="RIGHT", new_y="TOP")
    pdf.ln(7)
    pdf.set_font("DejaVu", "", 6.8)
    pdf.set_text_color(25, 39, 49)
    for index in range(half):
        pair = (
            records[index] if index < len(records) else None,
            records[index + half] if index + half < len(records) else None,
        )
        for item in pair:
            if item is None:
                values = ("", "", "", "")
            else:
                values = (
                    str(item["Grupo"]),
                    f"$ {item['Base propuesta / referencia (CUP)']:,.2f}",
                    f"{item['Factor aplicado']:.2f}×",
                    f"$ {item['Nueva propuesta (CUP)']:,.2f}",
                )
            for value, width in zip(values, widths[:4]):
                pdf.cell(width, 6.3, value, border=1, align="R" if value.startswith("$") else "C", new_x="RIGHT", new_y="TOP")
        pdf.ln(6.3)
    pdf.ln(2)
    pdf.set_font("DejaVu", "B", 7)
    pdf.set_text_color(130, 77, 12)
    pdf.multi_cell(0, 4.2, "Antes de usar o distribuir: verificar las fuentes normativas, resolver discrepancias y completar la revisión y autorización institucional correspondientes.")

    # Página 2: supuestos, incidencias detectadas y trazabilidad completa.
    pdf.add_page()
    pdf.set_y(13)
    pdf.set_font("DejaVu", "B", 13)
    pdf.set_text_color(21, 50, 75)
    pdf.cell(0, 8, "METODOLOGÍA, SUPUESTOS Y TRAZABILIDAD", align="C", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("DejaVu", "B", 8)
    pdf.set_text_color(130, 77, 12)
    pdf.cell(0, 6, f"{integrity_label} · versión {version}", align="C", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)
    heading("Supuestos y límites")
    for assumption in (
        "La propuesta se extrae de la plantilla asociada a la Resolución 87; los salarios oficiales y referencias normativas deben ser confirmados por el área responsable.",
        "Los grupos sin cargos propuestos no se multiplican: conservan la escala oficial de la Resolución 15/2026 (44 horas). El grupo I permanece en $ 3,210.00 CUP según la Resolución 14/2026.",
        f"El factor usado en este escenario es {factor:.2f}×. Cada resultado propuesto queda limitado a seis veces su salario fuente.",
        "No se calcula el impacto presupuestario total: no se dispone del conteo autorizado de plazas ocupadas/vacantes ni de componentes de nómina.",
        "Una huella SHA-256 coincidente detecta igualdad de bytes frente al manifiesto; no es firma digital ni prueba vigencia o validez jurídica.",
    ):
        paragraph(f"- {assumption}", size=6.8, line_height=3.9)
    anomalias_visibles = [
        _texto_visible(anomaly).strip()
        for anomaly in (anomalias or [])
        if not _es_vacio(anomaly) and _texto_visible(anomaly).strip()
    ]
    if anomalias_visibles:
        pdf.ln(1)
        heading("Incidencias observadas al extraer la fuente")
        for anomaly in anomalias_visibles[:4]:
            paragraph(f"- {anomaly[:320]}", size=6.4, line_height=3.6)
        if len(anomalias_visibles) > 4:
            paragraph(f"- Se omiten {len(anomalias_visibles) - 4} incidencias adicionales en esta página.", size=6.4, line_height=3.6)
    pdf.ln(2)
    heading("Fuentes supervisadas · hashes completos")
    checked_at = _texto_visible(manifiesto.get("checked_at_utc"), "No disponible")
    paragraph(
        f"Verificación UTC: {checked_at}    ·    "
        "Los hashes aprobados se comparan con el contenido actual del repositorio.",
        size=6.2,
        line_height=3.8,
    )
    for source in manifiesto.get("sources", []):
        source_id = _texto_visible(source.get("id"), "sin ID")
        label = _texto_visible(source.get("label"), "Fuente sin etiqueta")
        status = _texto_visible(source.get("status"), "NO VERIFICADA")
        path = _texto_visible(source.get("path"), "no disponible")
        expected_hash = _texto_visible(source.get("expected_sha256"), "no disponible")
        actual_hash = _texto_visible(source.get("actual_sha256"), "no disponible")
        pdf.set_font("DejaVu", "B", 6.2)
        pdf.set_text_color(21, 50, 75)
        pdf.cell(0, 4, f"{source_id} · {label} · {status}", new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("DejaVu", "", 5.0)
        pdf.set_text_color(35, 55, 70)
        pdf.multi_cell(0, 3.3, f"Ruta: {path}\nSHA-256 aprobado: {expected_hash}\nSHA-256 actual: {actual_hash}")
        pdf.ln(0.7)
    if manifiesto.get("error"):
        pdf.set_font("DejaVu", "B", 6.2)
        pdf.multi_cell(0, 3.8, f"Error de verificación: {str(manifiesto['error'])[:500]}")

    # Página 3: espacios en blanco para una decisión humana documentada.
    pdf.add_page()
    pdf.set_y(13)
    pdf.set_font("DejaVu", "B", 13)
    pdf.set_text_color(21, 50, 75)
    pdf.cell(0, 8, "REGISTRO DE REVISIÓN Y AUTORIZACIÓN", align="C", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("DejaVu", "B", 8)
    pdf.set_text_color(130, 77, 12)
    pdf.set_fill_color(255, 243, 220)
    pdf.cell(0, 7, "BORRADOR · NO APROBADO · COMPLETAR MANUALMENTE", border=1, align="C", fill=True, new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)
    paragraph(
        f"Escenario: factor {factor:.2f}× · Fecha del informe: {date.today():%d/%m/%Y} · "
        f"Versión de la aplicación: {version} · Integridad de fuentes: {integrity_label}",
        size=7.2,
        line_height=4.4,
    )
    pdf.ln(4)
    heading("Decisión (marcar una opción)")
    pdf.set_font("DejaVu", "", 8)
    pdf.set_text_color(35, 55, 70)
    pdf.cell(0, 7, "[  ] Aprobar propuesta     [  ] Devolver para correcciones     [  ] No aprobar", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(3)
    heading("Revisión mínima recomendada")
    paragraph(
        "[  ] Fuentes e integridad    [  ] Supuestos y metodología    [  ] Escala por grupo    "
        "[  ] Incidencias y normativa vigente",
        size=7.2,
        line_height=4.4,
    )
    pdf.ln(5)
    roles = ("PREPARÓ", "REVISÓ", "AUTORIZÓ")
    role_width = 60
    role_gap = 5
    signature_y = pdf.get_y()
    for index, role in enumerate(roles):
        x = 10 + index * (role_width + role_gap)
        y = signature_y
        pdf.rect(x, y, role_width, 45)
        pdf.set_xy(x + 2, y + 2)
        pdf.set_font("DejaVu", "B", 7.5)
        pdf.set_text_color(21, 50, 75)
        pdf.cell(role_width - 4, 5, role, new_x="LMARGIN", new_y="NEXT")
        pdf.set_x(x + 2)
        pdf.set_font("DejaVu", "", 6.5)
        pdf.set_text_color(35, 55, 70)
        pdf.multi_cell(role_width - 4, 7, "Nombre/cargo:\n\nFirma manuscrita:\n\nFecha:")
    pdf.set_y(signature_y + 49)
    heading("Observaciones, condiciones o referencias del acto de aprobación")
    y_notes = pdf.get_y() + 1
    pdf.rect(10, y_notes, 190, 54)
    pdf.set_y(y_notes + 57)
    pdf.set_font("DejaVu", "", 6.8)
    pdf.set_text_color(100, 110, 118)
    pdf.multi_cell(0, 4, "Las firmas deben ser realizadas por las personas facultadas conforme al procedimiento institucional. Este informe no contiene firmas electrónicas ni sustituye el acto administrativo que corresponda.")

    pdf.alias_nb_pages()
    pdf.set_auto_page_break(auto=False)
    for page_number in range(1, pdf.page_no() + 1):
        pdf.page = page_number
        pdf.set_y(-10)
        pdf.set_font("DejaVu", "", 6.5)
        pdf.set_text_color(100, 110, 118)
        pdf.cell(
            0,
            5,
            f"Capital Humano · v{version} · BORRADOR PARA REVISIÓN · Página {page_number} de {{nb}}",
            align="C",
        )
    return bytes(pdf.output())


# ---- Interfaz profesional / panel de control -------------------------------
NAV = [
    "Panel de control",
    "Escala salarial",
    "Comparar escenarios",
    "Plantilla por cargo",
    "Exportaciones",
    "Sistema y HTTPS",
    "Mejoras futuras",
]

st.markdown("""
<style>
:root{color-scheme:light;--ch-navy:#10234f;--ch-blue:#2251a2;--ch-cyan:#26b7d0;--ch-teal:#0f766e;--ch-ink:#19283c;--ch-muted:#586b7f;--ch-border:#dce5f2;--ch-canvas:#f1f5fc;--ch-rose:#b83748;--ch-amber:#9a5c08;--ch-violet:#7351cb}
html,body,[data-testid="stAppViewContainer"]{background:radial-gradient(ellipse at 92% 0%,rgba(86,164,255,.10),transparent 32rem),var(--ch-canvas);color:var(--ch-ink)}
.block-container{padding-top:1.25rem;padding-bottom:3.5rem;max-width:1540px}
[data-testid="stSidebar"]{background:linear-gradient(180deg,#111d46 0%,#0b1637 54%,#08122d 100%);border-right:1px solid rgba(157,185,255,.18);box-shadow:12px 0 34px rgba(16,35,79,.10)}
[data-testid="stSidebarContent"]{height:100vh;overflow-y:auto;overscroll-behavior:contain;scrollbar-gutter:stable}
[data-testid="stSidebar"] [data-testid="stMarkdownContainer"], [data-testid="stSidebar"] [data-testid="stWidgetLabel"], [data-testid="stSidebar"] label{color:#e5edff}
[data-testid="stSidebar"] [data-testid="stCaptionContainer"], [data-testid="stSidebar"] small{color:#b8c7e8!important}
[data-testid="stSidebar"] hr{border-color:rgba(194,213,255,.22)}
[data-testid="stSidebar"] [data-testid="stRadio"] label{border:1px solid transparent;border-radius:10px;padding:.35rem .55rem;transition:background .18s ease,border-color .18s ease}
[data-testid="stSidebar"] [data-testid="stRadio"] label:hover{background:rgba(103,153,255,.14);border-color:rgba(138,179,255,.24)}
[data-testid="stSidebar"] [data-testid="stRadio"] label:has(input:checked){background:linear-gradient(100deg,#2864d7,#397fe8);border-color:rgba(180,211,255,.48);box-shadow:0 5px 18px rgba(40,100,215,.26);color:#fff}
[data-testid="stSidebar"] [data-testid="stMetric"]{background:linear-gradient(145deg,#fff,#f3f7ff);border:1px solid rgba(196,214,247,.8);padding:.6rem .75rem;border-radius:13px}
[data-testid="stMetric"]{background:linear-gradient(145deg,rgba(255,255,255,.98),rgba(248,251,255,.96));border:1px solid var(--ch-border);border-radius:16px;padding:.9rem 1rem;box-shadow:0 8px 22px rgba(23,50,100,.055);min-height:100px;transition:transform .18s ease,box-shadow .18s ease}
[data-testid="stMetric"]:hover,.ch-kpi:hover,.ch-card:hover{transform:translateY(-2px);box-shadow:0 12px 28px rgba(35,75,145,.11)}
[data-testid="stMetricLabel"]{color:#52677b;font-weight:600}
[data-testid="stMetricValue"]{color:var(--ch-navy);font-variant-numeric:tabular-nums;font-weight:750}
.ch-hero{position:relative;overflow:hidden;background:linear-gradient(112deg,#fff 0%,#f5f8ff 54%,#eef7ff 100%);border:1px solid rgba(202,218,244,.94);border-left:5px solid var(--ch-cyan);padding:1.35rem 1.6rem;border-radius:18px;color:var(--ch-navy);margin-bottom:1.2rem;box-shadow:0 10px 30px rgba(23,50,100,.07)}
.ch-hero::after{content:"";position:absolute;right:-3rem;top:-5rem;width:16rem;height:16rem;border-radius:50%;background:radial-gradient(circle,rgba(38,183,208,.14),rgba(115,81,203,.05) 55%,transparent 70%);pointer-events:none}
.ch-hero h1{position:relative;color:var(--ch-navy);font-size:clamp(1.75rem,2.7vw,2.3rem);line-height:1.13;margin:0;font-weight:780;letter-spacing:-.025em}.ch-hero p{position:relative;color:#53677b;margin:.45rem 0 0;font-size:.98rem}
.ch-kicker{font-size:.72rem;letter-spacing:.13em;text-transform:uppercase;color:#087f8d;font-weight:780;margin-bottom:.36rem}
.ch-card{border:1px solid rgba(205,219,242,.92);background:linear-gradient(145deg,#fff,#f9fbff);border-radius:17px;padding:1rem 1.1rem;height:100%;box-shadow:0 7px 20px rgba(16,43,78,.045);transition:transform .18s ease,box-shadow .18s ease}
.ch-card h3{font-size:1.03rem;margin:.1rem 0 .45rem;color:var(--ch-navy)}
.ch-card p{margin:0;color:var(--ch-muted);font-size:.91rem;line-height:1.5}
.ch-muted{color:var(--ch-muted);font-size:.88rem}
.ch-status{display:inline-block;padding:.28rem .68rem;border-radius:999px;background:#e8f6f2;color:#176653;border:1px solid #ccebe1;font-size:.78rem;font-weight:750}
.ch-kpi{position:relative;overflow:hidden;display:flex;align-items:center;gap:.85rem;min-height:104px;padding:1rem 1.05rem;background:linear-gradient(140deg,#fff 0%,#f8faff 100%);border:1px solid rgba(203,218,244,.95);border-radius:17px;box-shadow:0 8px 22px rgba(23,50,100,.06);transition:transform .18s ease,box-shadow .18s ease}
.ch-kpi::before{content:"";position:absolute;inset:0 0 auto;height:3px;background:linear-gradient(90deg,#31b6dc,#496de2,#8a55d7)}
.ch-kpi-icon{width:50px;height:50px;flex:0 0 50px;border-radius:15px;display:grid;place-items:center;font-size:1.35rem;font-weight:800;color:#fff;box-shadow:0 7px 16px rgba(30,68,150,.2)}
.ch-kpi-icon.navy{background:linear-gradient(145deg,#345fc9,#263b91);color:#fff}.ch-kpi-icon.blue{background:linear-gradient(145deg,#24bed0,#147ea9);color:#fff}.ch-kpi-icon.teal{background:linear-gradient(145deg,#26b984,#14785f);color:#fff}.ch-kpi-icon.cyan{background:linear-gradient(145deg,#7959db,#4b43ad);color:#fff}.ch-kpi-icon.rose{background:linear-gradient(145deg,#df6278,#ae304d);color:#fff}
.ch-kpi-label{font-size:.78rem;font-weight:680;color:#586b7f;margin-bottom:.18rem}.ch-kpi-value{font-size:clamp(1.35rem,1.9vw,1.8rem);line-height:1.12;color:var(--ch-navy);font-weight:820;font-variant-numeric:tabular-nums;letter-spacing:-.02em}.ch-kpi-note{font-size:.76rem;color:#62758a;margin-top:.23rem}
.ch-section-title{font-size:1.04rem;font-weight:760;color:var(--ch-navy);margin:0 0 .65rem}
.ch-coverage-wrap{display:flex;align-items:center;justify-content:center;padding:.7rem 0 1rem}
.ch-donut{width:180px;height:180px;border-radius:50%;display:grid;place-items:center;position:relative}
.ch-donut::before{content:"";position:absolute;inset:20px;background:#fff;border-radius:50%}
.ch-donut-center{position:relative;text-align:center;color:var(--ch-navy);font-weight:800;font-size:1.6rem;line-height:1.2}.ch-donut-center small{display:block;font-size:.72rem;font-weight:600;color:#607286}
.ch-legend{display:flex;flex-wrap:wrap;justify-content:center;gap:.65rem 1rem;font-size:.78rem;color:#42566b}.ch-legend-item{display:flex;align-items:center;gap:.38rem}.ch-dot{width:10px;height:10px;border-radius:50%;display:inline-block}
.ch-summary-band{display:grid;grid-template-columns:repeat(3,minmax(0,1fr));gap:.5rem;padding:.9rem 1rem;background:linear-gradient(100deg,#e8f5fa,#f3f9fc);border:1px solid #d5eaf2;border-radius:14px;margin:.35rem 0 1.2rem;text-align:center}
.ch-summary-band strong{display:block;color:var(--ch-navy);font-size:1.13rem;font-variant-numeric:tabular-nums}.ch-summary-band span{font-size:.78rem;color:#526a7f}
.ch-priority{height:100%;padding:.9rem 1rem;border-radius:15px;background:linear-gradient(145deg,#fff,#f8faff);border:1px solid var(--ch-border);border-top:3px solid var(--ch-cyan);box-shadow:0 6px 16px rgba(23,50,100,.04)}.ch-priority.warn{border-top-color:#e2a334}.ch-priority.danger{border-top-color:#d4525f}.ch-priority.good{border-top-color:#1b9d80}
.ch-priority-head{font-size:.8rem;font-weight:750;color:#334b61;margin-bottom:.3rem}.ch-priority-value{font-size:1.03rem;font-weight:820;color:var(--ch-navy)}.ch-priority-note{font-size:.75rem;color:#65788a;margin-top:.18rem;line-height:1.4}
.ch-chart-toolbar{display:flex;align-items:center;justify-content:space-between;gap:.75rem;margin:.15rem 0 .35rem}.ch-chart-note{font-size:.76rem;color:#63758d;margin:.3rem 0 0}
[data-testid="stVerticalBlockBorderWrapper"]{background:linear-gradient(145deg,#fff,#fbfcff);border-color:var(--ch-border)!important;border-radius:18px;box-shadow:0 8px 24px rgba(16,43,78,.055)}
[data-testid="stDataFrame"]{border:1px solid var(--ch-border);border-radius:13px;overflow:hidden}
[data-testid="stSelectbox"] div[data-baseweb="select"]>div{border-radius:10px;border-color:#d6e0f0;background:rgba(255,255,255,.92)}
[data-testid="stTabs"] button[role="tab"]{border-radius:10px;font-weight:650}
.stButton button,.stDownloadButton button{border-radius:11px;font-weight:680;transition:transform .15s ease,box-shadow .15s ease}
.stButton button[kind="primary"],.stDownloadButton button[kind="primary"]{background:#123c67;border-color:#123c67;color:#fff}
.stButton button:hover,.stDownloadButton button:hover{border-color:#1592a5;color:#0c6074;box-shadow:0 4px 12px rgba(16,43,78,.12);transform:translateY(-1px)}
:focus-visible{outline:3px solid #087f9a!important;outline-offset:3px!important;box-shadow:0 0 0 2px #fff!important}
@media(max-width:760px){.block-container{padding:1rem .8rem 2rem}.ch-hero{padding:1rem 1.05rem;border-radius:13px}.ch-card{height:auto}.ch-summary-band{grid-template-columns:1fr}.ch-kpi{min-height:88px}.ch-donut{width:155px;height:155px}}
@media(max-width:640px){.block-container{padding:.85rem .65rem 1.75rem}.ch-hero h1{font-size:1.55rem}.ch-kpi{gap:.6rem;padding:.75rem}.ch-kpi-icon{width:42px;height:42px;flex-basis:42px}.ch-summary-band{margin:.25rem 0 1rem}}
@media(prefers-reduced-motion:reduce){*,*::before,*::after{scroll-behavior:auto!important;animation-duration:.01ms!important;animation-iteration-count:1!important;transition-duration:.01ms!important}}
iframe[title="Persistencia de posición de desplazamiento"]{position:fixed!important;left:-2px!important;bottom:0!important;width:1px!important;height:1px!important;border:0!important;opacity:0!important;pointer-events:none!important}
hr{border-color:var(--ch-border)}
</style>
""", unsafe_allow_html=True)

# Streamlit reruns after widget input. Keep only scroll offsets in browser
# sessionStorage so changing controls does not send or persist salary data.
st.iframe(
    """<script>
    (() => {
      try {
        const host = window.parent;
        const doc = host.document;
        const key = "capital-humano-scroll-v1";
        const main = doc.querySelector('[data-testid="stMain"]')
          || doc.querySelector('[data-testid="stAppViewContainer"]')
          || doc.scrollingElement;
        const sidebar = doc.querySelector('[data-testid="stSidebarContent"]')
          || doc.querySelector('[data-testid="stSidebar"]');
        const storage = host.sessionStorage;
        const save = () => {
          storage.setItem(key, JSON.stringify({
            main: Math.max(host.scrollY || 0, main?.scrollTop || 0),
            sidebar: sidebar?.scrollTop || 0
          }));
        };
        const restore = () => {
          const saved = JSON.parse(storage.getItem(key) || "{}");
          if (Number.isFinite(saved.main)) {
            if (main) main.scrollTop = saved.main;
            host.scrollTo(0, saved.main);
          }
          if (sidebar && Number.isFinite(saved.sidebar)) sidebar.scrollTop = saved.sidebar;
        };
        const bound = host.__capitalHumanoScrollNodes || (host.__capitalHumanoScrollNodes = new WeakSet());
        [main, sidebar].filter(Boolean).forEach(node => {
          if (!bound.has(node)) {
            node.addEventListener("scroll", save, { passive: true });
            bound.add(node);
          }
        });
        if (!host.__capitalHumanoScrollWindowBound) {
          host.addEventListener("scroll", save, { passive: true });
          host.__capitalHumanoScrollWindowBound = true;
        }
        requestAnimationFrame(restore);
        setTimeout(restore, 120);
      } catch (_) {
        // Some hosted browsers isolate components; native scrolling remains available.
      }
    })();
    </script>""",
    width="stretch",
    height=1,
    tab_index=-1,
    alt="Persistencia de posición de desplazamiento",
)

try:
    plantilla_sha256 = st.session_state.get("plantilla_importada", {}).get("sha256")
    propuesta, contadores, registros, anomalias, anterior, hallazgos_calidad = cargar_datos(plantilla_sha256)
except Exception as exc:
    st.error(T("No se pudo cargar la plantilla de origen: {error}", error=str(exc)))
    st.info(T("Comprueba que el archivo .xls de la plantilla y la escala histórica estén dentro del repositorio. Si las rutas son correctas, abre la guía de diagnóstico del README."))
    st.stop()

manifiesto = cargar_manifiesto()
version_visible = _texto_visible(manifiesto.get("application_version"), base.VERSION_APP)
fecha_verificacion_visible = _texto_visible(manifiesto.get("checked_at_utc"), "No disponible")
if st.session_state.get("plantilla_importada"):
    st.warning(
        T("Plantilla temporal activa solo en esta sesión. Es una fuente no aprobada: todas las descargas salariales están bloqueadas. Restaura la plantilla oficial para volver a exportar."),
        icon="⚠️",
    )
elif not manifiesto.get("integrity_ok"):
    st.warning(
        T("La copia actual de una o más fuentes o archivos de cálculo no coincide con el manifiesto aprobado. Revisa la sección Sistema y HTTPS y no distribuyas los informes hasta validar el cambio."),
        icon="⚠️",
    )

with st.sidebar:
    st.markdown(
        f"<div style='padding:.25rem 0 1rem'><div style='font-size:1.65rem'>▦ <b>Capital Humano</b></div><div style='color:#b8c7e8;font-size:.83rem'>{T('Panel de Escala Salarial')}</div></div>",
        unsafe_allow_html=True,
    )
    idioma_visible = st.selectbox("Idioma / Language", ["Español", "English"], key="idioma_interfaz")
    st.markdown(f"#### {T('Navegación')}")
    vistas_visibles = [T(item) for item in NAV]
    vista_elegida = st.radio(T("Secciones"), vistas_visibles, label_visibility="collapsed", key="panel_seccion")
    vista = NAV[vistas_visibles.index(vista_elegida)]
    st.divider()
    st.markdown(f"#### {T('Escenario salarial')}")
    factor = st.slider(
        T("Multiplicador común"), min_value=1.0, max_value=6.0,
        value=1.0, step=0.05, format="%.2f×", key="factor_global",
        help=T("Se aplica por cargo. El límite técnico nunca excede seis veces el salario de origen."),
    )
    st.metric(T("Tope permitido"), "6.00×")
    st.caption(T("El grupo I mantiene el mínimo oficial de $ 3,210.00 CUP."))
    st.divider()
    fuentes_etiqueta = (
        "Res. 14/2026 · Res. 15/2026 (44 h) · proposed template .xls"
        if UI_LANGUAGE == "en"
        else "Res. 14/2026 · Res. 15/2026 (44 h) · plantilla propuesta .xls"
    )
    st.caption(f"**{T('Fuentes activas')}**\n\n{fuentes_etiqueta}")

escala_df, detalle_df = crear_tablas(factor, propuesta, registros, anterior)
revision_calidad_base = crear_tabla_revision_calidad(hallazgos_calidad)
identidad_plantilla = st.session_state.get("plantilla_importada", {}).get("sha256", "plantilla-oficial")
firma_incidencias = hashlib.sha256(
    (identidad_plantilla + "|" + "|".join(revision_calidad_base["ID incidencia"].tolist())).encode("utf-8")
).hexdigest()
if st.session_state.get("firma_incidencias_calidad") != firma_incidencias:
    st.session_state["firma_incidencias_calidad"] = firma_incidencias
    st.session_state["revision_calidad_actual"] = revision_calidad_base.copy()
    st.session_state.pop("editor_revision_calidad", None)
assert detalle_df["Salario ajustado (CUP)"].le(detalle_df["Tope máximo 6× (CUP)"].round(2)).all()
propuesta_rows = escala_df[escala_df["Fuente del salario"] == "Plantilla Resolución 87"]


def encabezado(titulo, bajada):
    st.markdown(
        f"<div class='ch-hero'><div class='ch-kicker'>{T('Capital Humano · Gestión salarial')}</div><h1>{T(titulo)}</h1><p>{T(bajada)}</p></div>",
        unsafe_allow_html=True,
    )


def seleccionar_vista(nombre: str) -> None:
    """Callback para que las acciones del panel abran una sección real."""
    st.session_state["panel_seccion"] = T(nombre)


def mostrar_kpi(label, value, note, icon, tone):
    """Renderiza un KPI accesible; escapa todo el contenido interpolado en HTML."""
    accessible_label = html.escape(" · ".join(part for part in (label, value, note) if part))
    st.markdown(
        f"""<div class="ch-kpi" role="group" aria-label="{accessible_label}">
            <span class="ch-kpi-icon {html.escape(tone)}" aria-hidden="true">{html.escape(icon)}</span>
            <div><div class="ch-kpi-label">{html.escape(label)}</div>
            <div class="ch-kpi-value">{html.escape(value)}</div>
            <div class="ch-kpi-note">{html.escape(note)}</div></div>
        </div>""",
        unsafe_allow_html=True,
    )


def mostrar_estado_card(title, value, note, tone="good"):
    safe = [html.escape(text) for text in (title, value, note, tone)]
    st.markdown(
        f"""<div class="ch-priority {safe[3]}" role="group" aria-label="{safe[0]} · {safe[1]}">
            <div class="ch-priority-head">{safe[0]}</div>
            <div class="ch-priority-value">{safe[1]}</div>
            <div class="ch-priority-note">{safe[2]}</div>
        </div>""",
        unsafe_allow_html=True,
    )


def mostrar_dataframe(tabla, **kwargs):
    tabla_visual = ui.translate_dataframe(tabla, UI_LANGUAGE)
    column_config = kwargs.get("column_config")
    if column_config and UI_LANGUAGE == "en":
        kwargs["column_config"] = {T(label): config for label, config in column_config.items()}
    st.dataframe(tabla_visual, **kwargs)


TIPOS_GRAFICO = ["Barras horizontales", "Columnas", "Líneas", "Área"]
COLORES_GRAFICO = ["#4169d8", "#19aebd", "#8055d7", "#e4a02f"]


def mostrar_grafico(
    tabla: pd.DataFrame,
    tipo: str,
    *,
    currency: bool = True,
    height: int = 350,
) -> None:
    """Grafica series con valores y tooltips localizados, sin cambiar el dato."""
    if tabla.empty:
        st.info(T("No hay datos para mostrar con estos filtros."))
        return
    data = tabla.rename_axis("Grupo").reset_index().melt(
        id_vars="Grupo", var_name="Serie", value_name="Valor"
    )
    if currency:
        data["Etiqueta"] = data["Valor"].map(lambda value: ui.format_currency(value, UI_LANGUAGE))
        axis_expr = "datum.value < 0 ? '$ -' + format(-datum.value, ',.2f') : '$ ' + format(datum.value, ',.2f')"
        axis_title = T("Importe (CUP)")
        tooltip_title = T("Importe")
    else:
        data["Etiqueta"] = data["Valor"].map(lambda value: f"{value:,.2f}%")
        axis_expr = "format(datum.value, ',.1f') + '%'"
        axis_title = T("Variación porcentual")
        tooltip_title = T("Variación porcentual")

    palette = COLORES_GRAFICO[:max(1, data["Serie"].nunique())]
    color = alt.Color(
        "Serie:N",
        scale=alt.Scale(range=palette),
        legend=None if data["Serie"].nunique() == 1 else alt.Legend(title=None, orient="top"),
    )
    tooltip = [
        alt.Tooltip("Grupo:N", title=T("Grupo")),
        alt.Tooltip("Serie:N", title=T("Serie")),
        alt.Tooltip("Etiqueta:N", title=tooltip_title),
    ]
    if tipo == "Barras horizontales":
        chart = alt.Chart(data).mark_bar(cornerRadiusEnd=5, opacity=.94).encode(
            x=alt.X("Valor:Q", title=axis_title, axis=alt.Axis(labelExpr=axis_expr)),
            y=alt.Y("Grupo:N", title=None, sort=tabla.index.tolist()),
            color=color,
            tooltip=tooltip,
        )
    elif tipo == "Columnas":
        chart = alt.Chart(data).mark_bar(cornerRadiusTopLeft=4, cornerRadiusTopRight=4, opacity=.94).encode(
            x=alt.X("Grupo:N", title=None, sort=tabla.index.tolist(), axis=alt.Axis(labelAngle=-35)),
            y=alt.Y("Valor:Q", title=axis_title, axis=alt.Axis(labelExpr=axis_expr)),
            color=color,
            tooltip=tooltip,
        )
    elif tipo == "Líneas":
        chart = alt.Chart(data).mark_line(
            point=alt.OverlayMarkDef(filled=True, size=55), strokeWidth=3
        ).encode(
            x=alt.X("Grupo:N", title=None, sort=tabla.index.tolist()),
            y=alt.Y("Valor:Q", title=axis_title, axis=alt.Axis(labelExpr=axis_expr)),
            color=color,
            tooltip=tooltip,
        )
    else:
        chart = alt.Chart(data).mark_area(opacity=.24, line=True).encode(
            x=alt.X("Grupo:N", title=None, sort=tabla.index.tolist()),
            y=alt.Y("Valor:Q", title=axis_title, axis=alt.Axis(labelExpr=axis_expr)),
            color=color,
            tooltip=tooltip,
        )
    chart = (
        chart.properties(height=height)
        .configure_view(stroke=None)
        .configure_axis(labelColor="#5b6e86", titleColor="#435873", gridColor="#e8edf6")
        .configure_legend(labelColor="#435873")
        .interactive()
    )
    st.altair_chart(chart, width="stretch")


def mostrar_tabla_escala(tabla):
    mostrar_dataframe(
        tabla, width="stretch", hide_index=True, height=490,
        column_config={
            "Salario anterior Res. 29/2020 (CUP)": st.column_config.NumberColumn(format="$ %,.2f"),
            "Escala oficial Res. 15/2026 (CUP)": st.column_config.NumberColumn(format="$ %,.2f"),
            "Base propuesta / referencia (CUP)": st.column_config.NumberColumn(format="$ %,.2f"),
            "Tope máximo (6× base) (CUP)": st.column_config.NumberColumn(format="$ %,.2f"),
            "Nueva propuesta (CUP)": st.column_config.NumberColumn(format="$ %,.2f"),
            "Factor aplicado": st.column_config.NumberColumn(format="%.2f×"),
        },
    )


def mostrar_tabla_detalle(tabla):
    mostrar_dataframe(
        tabla, width="stretch", hide_index=True, height=560,
        column_config={
            "Id.": st.column_config.NumberColumn(format="%.0f"),
            "Salario original (CUP)": st.column_config.NumberColumn(format="$ %,.2f"),
            "Tope máximo 6× (CUP)": st.column_config.NumberColumn(format="$ %,.2f"),
            "Salario ajustado (CUP)": st.column_config.NumberColumn(format="$ %,.2f"),
            "Factor aplicado": st.column_config.NumberColumn(format="%.2f×"),
        },
    )


def mostrar_importador_plantilla():
    plantilla_temporal = st.session_state.get("plantilla_importada")
    if plantilla_temporal:
        if UI_LANGUAGE == "en":
            mensaje_activa = (
                f"Active in this session: {plantilla_temporal['filename']} · "
                f"{len(plantilla_temporal['records'])} valid rows · "
                f"SHA-256 {plantilla_temporal['sha256'][:16]}… Salary downloads are blocked."
            )
        else:
            mensaje_activa = (
                f"Activa en esta sesión: {plantilla_temporal['filename']} · "
                f"{len(plantilla_temporal['records'])} filas válidas · "
                f"SHA-256 {plantilla_temporal['sha256'][:16]}… Las descargas salariales están bloqueadas."
            )
        st.warning(mensaje_activa, icon="⚠️")
        if st.button(T("Volver a la plantilla oficial"), key="restaurar_plantilla_oficial"):
            st.session_state.pop("plantilla_importada", None)
            st.rerun()

    with st.expander(T("Importar otra plantilla para revisar"), expanded=False):
        st.caption(T(
            "Admite .xls y .xlsx sin macros (máximo 15 MiB). El archivo se analiza en memoria. "
            "La vista previa no lo activa ni modifica el archivo fuente."
        ))
        uploaded = st.file_uploader(
            T("Seleccionar plantilla (.xls o .xlsx)"),
            type=["xls", "xlsx"],
            key="subida_plantilla_calidad",
            help=T("No se guardan archivos subidos en el repositorio ni en el disco del servidor."),
        )
        if uploaded is None:
            return
        resultado_importacion = base.validar_plantilla_subida(uploaded.name, uploaded.getvalue())
        st.caption(T(
            "Archivo: {name} · {size} bytes · SHA-256: {digest}",
            name=resultado_importacion["filename"],
            size=f"{uploaded.size:,}",
            digest=resultado_importacion["sha256"],
        ))
        if resultado_importacion["validation"]:
            mostrar_dataframe(pd.DataFrame(resultado_importacion["validation"]), width="stretch", hide_index=True)
        if resultado_importacion["errors"]:
            for error in resultado_importacion["errors"]:
                st.error(T(error))
        if resultado_importacion["warnings"]:
            for warning in resultado_importacion["warnings"]:
                st.warning(T(warning))
        if resultado_importacion["preview"]:
            st.subheader(T("Vista previa · primeras 25 filas válidas"))
            mostrar_dataframe(pd.DataFrame(resultado_importacion["preview"]), width="stretch", hide_index=True)
        if resultado_importacion["findings"]:
            st.caption(T(
                "Hallazgos detectados: {count}. Se conservan todas las filas del archivo.",
                count=f"{len(resultado_importacion['findings']):,}",
            ))
        if resultado_importacion["ready"]:
            st.warning(T(
                "Confirmar solo cambia los datos de esta sesión. No reemplaza la fuente oficial; "
                "las descargas permanecen bloqueadas hasta aprobar formalmente la nueva fuente."
            ))
            if st.button(T("Usar temporalmente en esta sesión"), key="activar_plantilla_temporal"):
                st.session_state["plantilla_importada"] = {
                    key: resultado_importacion[key]
                    for key in (
                        "filename", "sha256", "size_bytes", "proposal", "counters",
                        "records", "anomalies", "findings",
                    )
                }
                st.rerun()


if vista == "Panel de control":
    encabezado("Panel de Control", "Resumen ejecutivo, validaciones y accesos directos para preparar la propuesta.")
    c1, c2, c3, c4 = st.columns(4, gap="medium")
    with c1:
        mostrar_kpi(T("Mínimo legal de referencia"), ui.format_currency(3210, UI_LANGUAGE), T("Grupo I · CUP"), "●", "navy")
    with c2:
        mostrar_kpi(T("Multiplicador seleccionado"), f"{factor:.2f}×", T("Límite absoluto 6× por cargo"), "▥", "blue")
    with c3:
        mostrar_kpi(T("Cargos encontrados"), f"{len(detalle_df):,}", T("Registros de la plantilla fuente"), "▤", "teal")
    with c4:
        mostrar_kpi(T("Grupos con propuesta"), f"{len(propuesta)}/{len(escala_df)}", T("Los demás conservan la escala oficial"), "✓", "cyan")

    left, right = st.columns([1.7, 1], gap="large")
    with left:
        with st.container(border=True):
            st.markdown(f"<div class='ch-section-title'>{T('Vista comparativa de la escala')}</div>", unsafe_allow_html=True)
            controls = st.columns([1.25, 1], gap="small")
            desglose = controls[0].selectbox(
                T("Desglose del gráfico"),
                ["Comparación oficial / propuesta", "Variación absoluta (CUP)", "Variación porcentual"],
                format_func=T,
                key="dashboard_chart_breakdown",
            )
            tipo_grafico = controls[1].selectbox(
                T("Tipo de gráfico"), TIPOS_GRAFICO, format_func=T,
                key="dashboard_chart_type",
            )
            if desglose == "Comparación oficial / propuesta":
                grafico = propuesta_rows[[
                    "Grupo", "Escala oficial Res. 15/2026 (CUP)", "Nueva propuesta (CUP)"
                ]].set_index("Grupo").rename(columns={
                    "Escala oficial Res. 15/2026 (CUP)": T("Oficial Res. 15/2026"),
                    "Nueva propuesta (CUP)": f"{T('Propuesta')} {factor:.2f}×",
                })
            elif desglose == "Variación porcentual":
                grafico = propuesta_rows.assign(
                    **{"Cambio porcentual": (
                        (propuesta_rows["Nueva propuesta (CUP)"] /
                         propuesta_rows["Escala oficial Res. 15/2026 (CUP)"] - 1) * 100
                    )}
                ).set_index("Grupo")[["Cambio porcentual"]].rename(
                    columns={"Cambio porcentual": T("Variación porcentual")}
                )
            else:
                grafico = propuesta_rows.assign(
                    **{"Cambio (CUP)": (
                        propuesta_rows["Nueva propuesta (CUP)"] -
                        propuesta_rows["Escala oficial Res. 15/2026 (CUP)"]
                    )}
                ).set_index("Grupo")[["Cambio (CUP)"]].rename(
                    columns={"Cambio (CUP)": T("Variación absoluta (CUP)")}
                )
            mostrar_grafico(
                grafico, tipo_grafico,
                currency=desglose != "Variación porcentual",
                height=345,
            )
            st.caption(T("Los datos mostrados corresponden a grupos con propuesta; no representan el costo total de nómina."))
    with right:
        with st.container(border=True):
            st.markdown(f"<div class='ch-section-title'>{T('Cobertura de la escala')}</div>", unsafe_allow_html=True)
            grupos_con_propuesta = len(propuesta)
            total_grupos = len(escala_df)
            grupos_oficiales = max(0, total_grupos - grupos_con_propuesta)
            cobertura = grupos_con_propuesta / total_grupos if total_grupos else 0
            giro = cobertura * 360
            texto_accesible = T(
                "{proposed} de {total} grupos tienen propuesta de plantilla; {official} conservan la escala oficial.",
                proposed=grupos_con_propuesta,
                total=total_grupos,
                official=grupos_oficiales,
            )
            st.markdown(
                f"<div class='ch-coverage-wrap'><div class='ch-donut' role='img' aria-label='{html.escape(texto_accesible)}' style='background:conic-gradient(#173b65 0deg {giro:.1f}deg,#35b9d0 {giro:.1f}deg 360deg)'>"
                f"<div class='ch-donut-center'>{grupos_con_propuesta}/{total_grupos}<small>{T('Grupos con propuesta')}</small></div></div></div>",
                unsafe_allow_html=True,
            )
            st.markdown(
                f"<div class='ch-legend'><span class='ch-legend-item'><i class='ch-dot' style='background:#173b65'></i>{T('Con propuesta')}</span>"
                f"<span class='ch-legend-item'><i class='ch-dot' style='background:#35b9d0'></i>{T('Escala oficial conservada')}</span></div>",
                unsafe_allow_html=True,
            )
    st.markdown(
        f"<div class='ch-summary-band'><div><strong>{total_grupos}</strong><span>{T('Grupos en la escala')}</span></div>"
        f"<div><strong>{grupos_con_propuesta}</strong><span>{T('Con propuesta de plantilla')}</span></div>"
        f"<div><strong>{grupos_oficiales}</strong><span>{T('Con escala oficial conservada')}</span></div></div>",
        unsafe_allow_html=True,
    )
    st.caption(T("No se estima el gasto total de nómina: esta aplicación no dispone de un registro autorizado de plazas ocupadas y vacantes."))
    st.divider()
    st.subheader(T("Atención prioritaria"))
    p1, p2, p3 = st.columns(3, gap="medium")
    with p1:
        integrity_title = T("Integridad de fuentes")
        integrity_value = T("Verificada" if manifiesto.get("integrity_ok") else "Revisar")
        mostrar_estado_card(integrity_title, integrity_value, T("Las fuentes se comparan con las huellas SHA-256 aprobadas."), "good" if manifiesto.get("integrity_ok") else "danger")
    with p2:
        incidencias = len(hallazgos_calidad)
        mostrar_estado_card(
            T("Calidad de la plantilla"),
            T("Sin incidencias" if incidencias == 0 else "{count} por revisar", count=incidencias),
            T("Los hallazgos no se corrigen ni eliminan automáticamente."),
            "good" if incidencias == 0 else "warn",
        )
    with p3:
        mostrar_estado_card(
            T("Impacto total de nómina"),
            T("No estimado"),
            T("Requiere datos autorizados de plazas y validación de Finanzas."),
            "warn",
        )
    st.divider()
    st.subheader(T("Acciones rápidas"))
    a1, a2, a3 = st.columns(3, gap="medium")
    with a1:
        st.markdown(f"<div class='ch-card'><h3>{T('📊 Revisar escala')}</h3><p>{T('Consulta los 32 grupos, salarios oficiales y propuesta ajustada.')}</p></div>", unsafe_allow_html=True)
        st.button(T("Abrir escala"), key="quick_open_scale", on_click=seleccionar_vista, args=("Escala salarial",), width="stretch")
    with a2:
        st.markdown(f"<div class='ch-card'><h3>{T('👥 Auditar plantilla')}</h3><p>{T('Id. correlativo, búsqueda y filtros conservan la referencia a hoja y fila de origen.')}</p></div>", unsafe_allow_html=True)
        st.button(T("Abrir plantilla"), key="quick_open_template", on_click=seleccionar_vista, args=("Plantilla por cargo",), width="stretch")
    with a3:
        st.markdown(f"<div class='ch-card'><h3>{T('📄 Preparar entregables')}</h3><p>{T('Exporta el escenario activo a un libro Excel o PDF imprimible.')}</p></div>", unsafe_allow_html=True)
        st.button(T("Abrir exportaciones"), key="quick_open_exports", on_click=seleccionar_vista, args=("Exportaciones",), width="stretch")
    if anomalias:
        with st.expander(T("Observaciones detectadas en la fuente"), expanded=False):
            for nota in anomalias:
                st.warning(T(nota))

elif vista == "Escala salarial":
    encabezado("Escala Salarial", T("Escenario en vivo · multiplicador {factor}× · límite absoluto 6× de la base de origen.", factor=f"{factor:.2f}"))
    st.info(T("Los grupos con propuesta se recalculan desde los salarios del archivo de plantilla. Los grupos sin propuesta conservan la escala oficial de la Resolución 15/2026; el grupo I permanece en $ 3,210.00."))
    mostrar_tabla_escala(escala_df)
    with st.expander(T("Metodología y supuestos")):
        if UI_LANGUAGE == "en":
            st.markdown("""
            - The global multiplier is applied to each position's source salary.
            - The six-times cap is checked for every record.
            - The consolidated group table uses the majority value when the .xls contains discrepancies.
            - The detailed table retains every individual amount for audit.
            - No values are invented for groups absent from the template.
            """)
        else:
            st.markdown("""
            - El multiplicador global se aplica al salario por cargo de la plantilla fuente.
            - El tope de seis veces se comprueba individualmente en cada registro.
            - Para la tabla consolidada de grupo se usa el valor mayoritario cuando el .xls contiene discrepancias.
            - La tabla detallada conserva cada importe individual para auditoría.
            - Los grupos ausentes en la plantilla no reciben importes inventados.
            """)

elif vista == "Comparar escenarios":
    encabezado("Comparar escenarios", "Contrasta dos factores en paralelo por grupo o por cargo, sin convertir la suma en un presupuesto de nómina.")
    ca, cb = st.columns(2)
    factor_a = ca.slider(T("Escenario A · factor"), 1.0, 6.0, 1.0, 0.05, format="%.2f×", key="comparador_a")
    factor_b = cb.slider(T("Escenario B · factor"), 1.0, 6.0, 1.5, 0.05, format="%.2f×", key="comparador_b")
    grupo_tab, cargo_tab = st.tabs([T("Resumen por grupo"), T("Detalle por cargo")])

    with grupo_tab:
        comparacion_df = crear_comparacion_escenarios(factor_a, factor_b, propuesta, anterior)
        grupos_modificados = comparacion_df["Variación B − A (CUP)"].abs().gt(0.005).sum()
        maxima_variacion = comparacion_df["Variación B − A (CUP)"].abs().max()
        k1, k2, k3 = st.columns(3)
        k1.metric(T("Grupos con cambio"), f"{grupos_modificados} {T('de')} {len(comparacion_df)}")
        k2.metric(T("Mayor cambio por grupo"), ui.format_currency(maxima_variacion, UI_LANGUAGE))
        k3.metric(T("Factores comparados"), f"{factor_a:.2f}× {T('vs.')} {factor_b:.2f}×")
        chart_controls = st.columns([1.25, 1], gap="small")
        desglose_comparacion = chart_controls[0].selectbox(
            T("Desglose del gráfico"),
            ["Variación absoluta (CUP)", "Variación porcentual", "Comparación escenarios A/B"],
            format_func=T,
            key="comparison_chart_breakdown",
        )
        tipo_grafico_comparacion = chart_controls[1].selectbox(
            T("Tipo de gráfico"), TIPOS_GRAFICO, format_func=T,
            key="comparison_chart_type",
        )
        if desglose_comparacion == "Comparación escenarios A/B":
            grafico_comparacion = comparacion_df.set_index("Grupo")[[
                "Escenario A (CUP)", "Escenario B (CUP)"
            ]].rename(columns={
                "Escenario A (CUP)": T("Escenario A"),
                "Escenario B (CUP)": T("Escenario B"),
            })
        elif desglose_comparacion == "Variación porcentual":
            grafico_comparacion = comparacion_df.assign(
                **{"Variación porcentual": comparacion_df["Variación relativa"] * 100}
            ).set_index("Grupo")[["Variación porcentual"]].rename(
                columns={"Variación porcentual": T("Variación porcentual")}
            )
        else:
            grafico_comparacion = comparacion_df.set_index("Grupo")[["Variación B − A (CUP)"]].rename(
                columns={"Variación B − A (CUP)": T("Variación absoluta (CUP)")}
            )
        mostrar_grafico(
            grafico_comparacion,
            tipo_grafico_comparacion,
            currency=desglose_comparacion != "Variación porcentual",
            height=350,
        )
        mostrar_dataframe(
            comparacion_df, width="stretch", hide_index=True,
            column_config={
                "Escenario A (CUP)": st.column_config.NumberColumn(format="$ %,.2f"),
                "Escenario B (CUP)": st.column_config.NumberColumn(format="$ %,.2f"),
                "Variación B − A (CUP)": st.column_config.NumberColumn(format="$ %,.2f"),
                "Variación relativa": st.column_config.NumberColumn(format="percent"),
            },
        )
        st.warning(T("La suma de importes por grupo no equivale al impacto total de nómina: no se dispone aquí de la cantidad de plazas ocupadas/vacantes por grupo. No uses esta comparación como presupuesto de nómina."))
        out_csv, out_xlsx = st.columns(2)
        csv_comparacion = crear_zip_csv_con_manifiesto(comparacion_df, "Comparacion_escenarios_grupo.csv", manifiesto)
        with out_csv:
            st.download_button(T("Descargar CSV + manifiesto (.zip)"), csv_comparacion, "Comparacion_escenarios_grupo.zip", "application/zip", width="stretch", disabled=descargas_deshabilitadas())
        with out_xlsx:
            xlsx_comparacion = crear_excel_comparacion(comparacion_df, factor_a, factor_b, manifiesto)
            st.download_button(
                T("Descargar comparación Excel"), xlsx_comparacion,
                f"Comparacion_escenarios_grupo_{factor_a:.2f}x_{factor_b:.2f}x.xlsx",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                width="stretch",
                disabled=descargas_deshabilitadas(),
            )

    with cargo_tab:
        comparacion_cargos_df = crear_comparacion_cargos(factor_a, factor_b, registros)
        cargos_modificados = comparacion_cargos_df["Variación B − A (CUP)"].abs().gt(0.005).sum()
        maximo_cambio_cargo = comparacion_cargos_df["Variación B − A (CUP)"].abs().max() if not comparacion_cargos_df.empty else 0
        d1, d2, d3 = st.columns(3)
        d1.metric(T("Cargos con cambio"), f"{cargos_modificados} {T('de')} {len(comparacion_cargos_df)}")
        d2.metric(T("Mayor cambio por cargo"), ui.format_currency(maximo_cambio_cargo, UI_LANGUAGE))
        d3.metric(T("Factores comparados"), f"{factor_a:.2f}× {T('vs.')} {factor_b:.2f}×")
        mostrar_dataframe(
            comparacion_cargos_df, width="stretch", hide_index=True,
            column_config={
                "Id.": st.column_config.NumberColumn(format="%.0f"),
                "Fila de origen": st.column_config.NumberColumn(format="%.0f"),
                "Salario fuente (CUP)": st.column_config.NumberColumn(format="$ %,.2f"),
                "Salario A (CUP)": st.column_config.NumberColumn(format="$ %,.2f"),
                "Salario B (CUP)": st.column_config.NumberColumn(format="$ %,.2f"),
                "Factor A": st.column_config.NumberColumn(format="%.2f×"),
                "Factor B": st.column_config.NumberColumn(format="%.2f×"),
                "Variación B − A (CUP)": st.column_config.NumberColumn(format="$ %,.2f"),
                "Variación relativa": st.column_config.NumberColumn(format="percent"),
            },
        )
        st.warning(T("Esta comparación es por cargo y no calcula costo total de nómina. Los dos salarios de cada fila respetan el tope absoluto de 6× la fuente."))
        detalle_csv, detalle_xlsx = st.columns(2)
        paquete_detalle_csv = crear_zip_csv_con_manifiesto(
            comparacion_cargos_df, "Comparacion_escenarios_por_cargo.csv", manifiesto
        )
        with detalle_csv:
            st.download_button(
                T("Descargar CSV por cargo + manifiesto (.zip)"), paquete_detalle_csv,
                "Comparacion_escenarios_por_cargo.zip", "application/zip", width="stretch",
                disabled=descargas_deshabilitadas(),
            )
        with detalle_xlsx:
            detalle_xlsx_bytes = crear_excel_comparacion_cargos(
                comparacion_cargos_df, factor_a, factor_b, manifiesto
            )
            st.download_button(
                T("Descargar comparación por cargo Excel"), detalle_xlsx_bytes,
                f"Comparacion_escenarios_cargos_{factor_a:.2f}x_{factor_b:.2f}x.xlsx",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                width="stretch",
                disabled=descargas_deshabilitadas(),
            )

elif vista == "Plantilla por cargo":
    encabezado("Plantilla por cargo", "Detalle rastreable al archivo de origen con filtros y búsqueda rápida.")
    mostrar_importador_plantilla()
    st.subheader(T("Calidad de la fuente"))
    hallazgos_df = pd.DataFrame(hallazgos_calidad)
    conteo_hallazgos = (
        hallazgos_df["Tipo de incidencia"].value_counts().to_dict()
        if not hallazgos_df.empty else {}
    )
    q1, q2, q3, q4, q5 = st.columns(5)
    q1.metric(T("Registros válidos"), f"{len(registros):,}")
    q2.metric(T("Salarios ausentes"), int(conteo_hallazgos.get("Salario ausente", 0)))
    q3.metric(T("Salarios inválidos"), int(conteo_hallazgos.get("Salario inválido", 0)))
    q4.metric(T("Posibles duplicados"), int(conteo_hallazgos.get("Posible duplicado", 0)))
    q5.metric(T("Discrepancias"), int(conteo_hallazgos.get("Discrepancia salarial de grupo", 0)))
    if hallazgos_df.empty:
        st.success(T("No se detectaron incidencias en las filas con grupo válido. Esto no sustituye la revisión humana de la fuente."))
    else:
        st.warning(T(
            "Se detectaron {count} incidencias para revisar. Este diagnóstico no corrige, elimina ni modifica filas de la fuente. La escala consolidada mantiene la regla actual del valor mayoritario ante discrepancias.",
            count=f"{len(hallazgos_df):,}",
        ))
        mostrar_dataframe(hallazgos_df, width="stretch", hide_index=True)
        st.subheader(T("Revisión manual de incidencias"))
        st.caption(T("El estado y la nota se conservan durante esta sesión y en la descarga; no se escriben en la plantilla original ni en los archivos fuente."))
        revision_para_editar = ui.translate_dataframe(st.session_state["revision_calidad_actual"], UI_LANGUAGE)
        campos_bloqueados = [
            "ID incidencia", "Tipo de incidencia", "Hoja de origen", "Fila de origen",
            "Cargo", "Grupo", "Valor detectado", "Detalle",
        ]
        opciones_estado = ["Pendiente", "En revisión", "Revisada", "No aplica"]
        revision_editada = st.data_editor(
            revision_para_editar,
            width="stretch", hide_index=True, key=f"editor_revision_calidad_{UI_LANGUAGE}",
            disabled=[T(campo) for campo in campos_bloqueados],
            column_config={
                T("Estado de revisión"): st.column_config.SelectboxColumn(
                    T("Estado de revisión"),
                    options=[T(estado) for estado in opciones_estado],
                    required=True,
                ),
                T("Nota de revisión"): st.column_config.TextColumn(
                    T("Nota de revisión"), max_chars=500,
                ),
            },
        )
        revision_calidad_df = revision_editada
        if UI_LANGUAGE == "en":
            columnas_origen = st.session_state["revision_calidad_actual"].columns
            nombres_es = {T(columna): columna for columna in columnas_origen}
            revision_calidad_df = revision_calidad_df.rename(columns=nombres_es)
            estados_es = {T(estado): estado for estado in opciones_estado}
            revision_calidad_df["Estado de revisión"] = revision_calidad_df["Estado de revisión"].map(estados_es).fillna("Pendiente")
        st.session_state["revision_calidad_actual"] = revision_calidad_df.copy()
        paquete_calidad = crear_zip_csv_con_manifiesto(
            revision_calidad_df, "Incidencias_calidad_revisadas.csv", manifiesto
        )
        st.download_button(
            T("Descargar incidencias + manifiesto (.zip)"), paquete_calidad,
            "Incidencias_calidad_plantilla.zip", "application/zip",
            key="descargar_incidencias_calidad",
            disabled=descargas_deshabilitadas(),
        )
    st.caption(T("Cada hallazgo incluye hoja y fila. Un posible duplicado coincide en hoja, grupo, cargo y salario: puede ser una plaza legítima repetida, por lo que se conserva hasta revisión humana."))
    st.divider()
    f1, f2, f3 = st.columns([1, 1, 2])
    grupos = sorted(detalle_df["Grupo propuesto"].unique(), key=lambda g: base.ROMANO_A_ENTERO[g])
    hojas = sorted(detalle_df["Hoja de origen"].unique())
    filtro_grupo = f1.multiselect(T("Grupo"), grupos, placeholder=T("Todos los grupos"))
    filtro_hoja = f2.multiselect(T("Hoja fuente"), hojas, placeholder=T("Todas las hojas"))
    busqueda = f3.text_input(T("Buscar cargo"), placeholder=T("Ej.: director, técnico, operador…"))
    filtrada = detalle_df.copy()
    if filtro_grupo:
        filtrada = filtrada[filtrada["Grupo propuesto"].isin(filtro_grupo)]
    if filtro_hoja:
        filtrada = filtrada[filtrada["Hoja de origen"].isin(filtro_hoja)]
    if busqueda.strip():
        filtrada = filtrada[filtrada["Cargo"].str.contains(busqueda.strip(), case=False, regex=False, na=False)]
    m1, m2, m3 = st.columns(3)
    m1.metric(T("Registros visibles"), f"{len(filtrada):,}")
    m2.metric(T("Salario fuente total"), ui.format_currency(filtrada["Salario original (CUP)"].sum(), UI_LANGUAGE))
    m3.metric(T("Propuesta calculada"), ui.format_currency(filtrada["Salario ajustado (CUP)"].sum(), UI_LANGUAGE))
    mostrar_tabla_detalle(filtrada)
    csv = crear_zip_csv_con_manifiesto(filtrada, "Plantilla_Resolucion87_filtrada.csv", manifiesto)
    st.download_button(
        T("Descargar filtro CSV + manifiesto (.zip)"), csv,
        "Plantilla_Resolucion87_filtrada.zip", "application/zip",
        disabled=descargas_deshabilitadas(),
    )

elif vista == "Exportaciones":
    encabezado("Exportar y presentar", "Valida el escenario antes de preparar documentos consistentes para revisión.")
    st.write(f"**{T('Escenario activo')}:** {factor:.2f}× · **{T('Cargos')}:** {len(detalle_df)} · **{T('Fecha')}:** {ui.format_date(date.today(), UI_LANGUAGE)}")
    st.info(T("Excel, PDF e informe de aprobación se generan desde la misma escala, detalle, factor y manifiesto. El informe de aprobación permanece como borrador."))
    preflight_df, exportacion_lista = validar_preexportacion(
        escala_df, detalle_df, factor, manifiesto, len(registros)
    )
    st.subheader(T("Comprobaciones previas a la descarga"))
    mostrar_dataframe(preflight_df, width="stretch", hide_index=True)
    if not exportacion_lista:
        st.error(T("Las descargas quedan bloqueadas hasta corregir las comprobaciones marcadas como ERROR. No se modificó ninguna fuente."))
    else:
        if hallazgos_calidad:
            st.warning(T(
                "Hay {count} incidencias de calidad pendientes de revisión. Puedes descargar los informes para revisión, pero no interpretes su generación como resolución de esas incidencias.",
                count=f"{len(hallazgos_calidad):,}",
            ))
        with st.spinner(T("Preparando los archivos para descarga…")):
            excel_bytes = crear_excel(escala_df, detalle_df, factor, manifiesto)
            pdf_bytes = crear_pdf(escala_df, detalle_df, factor, manifiesto)
            informe_aprobacion_bytes = crear_informe_aprobacion_pdf(
                escala_df, detalle_df, factor, manifiesto, anomalias
            )
        e1, e2, e3 = st.columns(3)
        with e1:
            st.markdown(f"<div class='ch-card'><h3>{T('📗 Libro Excel')}</h3><p>{T('32 grupos, formatos y detalle completo de la plantilla fuente.')}</p></div>", unsafe_allow_html=True)
            st.download_button(
                T("Descargar Excel (.xlsx)"), excel_bytes,
                f"Propuesta_Escala_Resolucion87_{factor:.2f}x.xlsx",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                width="stretch",
            )
        with e2:
            st.markdown(f"<div class='ch-card'><h3>{T('📕 Informe PDF')}</h3><p>{T('Escala resumida, detalle paginado y página de trazabilidad.')}</p></div>", unsafe_allow_html=True)
            st.download_button(
                T("Descargar PDF imprimible"), pdf_bytes,
                f"Propuesta_Escala_Resolucion87_{factor:.2f}x.pdf",
                "application/pdf", width="stretch",
            )
        with e3:
            st.markdown(f"<div class='ch-card'><h3>{T('📝 Informe para aprobación')}</h3><p>{T('Metodología, resumen por grupo, fuentes y espacios para decisión y firmas manuscritas.')}</p></div>", unsafe_allow_html=True)
            st.download_button(
                T("Descargar borrador de aprobación"),
                informe_aprobacion_bytes,
                f"Informe_Aprobacion_Capital_Humano_{factor:.2f}x.pdf",
                "application/pdf", width="stretch",
            )
        paquete_completo = crear_paquete_completo(
            excel_bytes,
            pdf_bytes,
            informe_aprobacion_bytes,
            factor,
            detalle_df,
            st.session_state.get("revision_calidad_actual", revision_calidad_base),
            manifiesto,
        )
        st.subheader(T("Paquete completo para revisión"))
        st.caption(T("Incluye los tres documentos, parámetros del escenario, incidencias con revisión manual, manifiesto de fuentes y un inventario SHA-256 de todos los demás archivos del paquete."))
        st.download_button(
            T("Descargar paquete completo (.zip)"),
            paquete_completo,
            f"Paquete_Capital_Humano_{factor:.2f}x.zip",
            "application/zip",
            width="stretch",
            key="descargar_paquete_completo",
        )
        st.caption(T("Versión {version} · integridad verificada · el informe de aprobación permanece marcado como borrador.", version=version_visible))

elif vista == "Sistema y HTTPS":
    encabezado("Sistema y seguridad", "Comprobaciones de las fuentes, del entorno de ejecución y del canal seguro.")
    ssl_cert = str(st.get_option("server.sslCertFile") or "").strip()
    try:
        request_headers = {str(k).lower(): str(v) for k, v in st.context.headers.items()}
    except Exception:
        request_headers = {}
    proxy_tls = request_headers.get("x-forwarded-proto", "").lower() == "https"
    cert_file_ok = bool(ssl_cert and Path(ssl_cert).is_file())
    tls_state = T("HTTPS confirmado por proxy") if proxy_tls else (T("Certificado TLS nativo configurado") if cert_file_ok else T("No confirmado por esta sesión"))
    q1, q2, q3, q4 = st.columns(4)
    q1.metric(T("Canal"), tls_state)
    q2.metric(T("Plantilla fuente"), T("Disponible" if base.RUTA_PLANTILLA.is_file() else "No encontrada"))
    q3.metric(T("Escala histórica"), T("Disponible" if base.RUTA_ESCALA_BASE.is_file() else "No encontrada"))
    q4.metric(T("Integridad SHA-256"), T("Verificada" if manifiesto.get("integrity_ok") else "Revisar"))
    st.markdown(f"#### {T('Información de ejecución')}")
    sys_df = pd.DataFrame([
        {"Comprobación": "Versión de Capital Humano", "Resultado": version_visible},
        {"Comprobación": "Integridad de fuentes", "Resultado": "Verificada" if manifiesto.get("integrity_ok") else "Revisión requerida"},
        {"Comprobación": "Verificación de fuentes (UTC)", "Resultado": fecha_verificacion_visible},
        {"Comprobación": "Sistema operativo", "Resultado": platform.platform()},
        {"Comprobación": "Python", "Resultado": sys.version.split()[0]},
        {"Comprobación": "Directorio de la aplicación", "Resultado": str(CARPETA)},
        {"Comprobación": "Configuración de certificado nativo", "Resultado": "Presente" if cert_file_ok else "No detectada"},
        {"Comprobación": "Proxy HTTPS", "Resultado": "Detectado" if proxy_tls else "No detectado"},
        {"Comprobación": "Validación de tope 6×", "Resultado": "Correcta" if detalle_df["Salario ajustado (CUP)"].le(detalle_df["Tope máximo 6× (CUP)"].round(2)).all() else "Revisar"},
    ])
    mostrar_dataframe(sys_df, width="stretch", hide_index=True)
    st.markdown(f"#### {T('Huellas SHA-256 de fuentes y código')}")
    if manifiesto.get("integrity_ok"):
        st.success(T("Todos los archivos supervisados coinciden byte a byte con las huellas aprobadas del manifiesto."))
    else:
        st.error(T("Una fuente, el manifiesto o el código no coincide con la versión registrada; valida el cambio antes de distribuir informes."))
        if manifiesto.get("error"):
            st.caption(manifiesto["error"])
    manifest_df = pd.DataFrame(manifiesto.get("sources", []))
    if not manifest_df.empty:
        mostrar_dataframe(
            manifest_df, width="stretch", hide_index=True,
            column_config={
                "expected_sha256": st.column_config.TextColumn(T("SHA-256 aprobado"), width="large"),
                "actual_sha256": st.column_config.TextColumn(T("SHA-256 actual"), width="large"),
                "size_bytes": st.column_config.NumberColumn(T("Tamaño (bytes)"), format="localized"),
            },
        )
    st.caption(T("Para actualizar el manifiesto tras un cambio autorizado, revisa primero la nueva fuente y ejecuta: python tools/actualizar_manifiesto_fuentes.py. El hash no sustituye revisión normativa ni firma digital."))
    st.warning(T("HTTPS cifra y autentica el servidor; no sustituye controles de acceso a nivel de aplicación. Para uso público, usa Caddy + Let's Encrypt y autenticación Basic configurada en el despliegue. No expongas directamente el puerto 8501."))
    if os.name == "nt":
        control_script = CARPETA / "control_app.ps1"
        st.markdown(f"#### {T('Inicio automático HTTPS en Windows')}")
        if control_script.is_file():
            st.success(T("Está disponible el controlador de inicio/estado de la app."))
            st.code(".\\control_app.ps1 -Accion instalar-autoinicio\n.\\control_app.ps1 -Accion estado\n.\\control_app.ps1 -Accion detener\n.\\control_app.ps1 -Accion quitar-autoinicio", language="powershell")
            st.caption(T("El inicio automático se programa al iniciar sesión del mismo usuario; la CA mkcert debe instalarse una vez en ese perfil."))
        else:
            st.info(T("El control de inicio automático se habilita con el script de gestión de Windows incluido en esta versión."))
    else:
        st.markdown(f"#### {T('Autoarranque del servidor')}")
        st.info(T("En Windows usa el programador de tareas del script de control. En servidores Linux, se recomienda el despliegue Docker Compose con Caddy y políticas de reinicio; certificados persistentes en el volumen caddy_data."))
    with st.expander(T("Recomendaciones operativas")):
        if UI_LANGUAGE == "en":
            st.markdown("""
            - Update Python and dependencies in an x64 virtual environment (Windows).
            - Do not share `.certs`, private keys, `deploy/.env`, or credentials.
            - In production, expose only ports 80/443 through Caddy; do not expose 8501 to the Internet.
            - Check HTTPS in the browser and confirm the certificate date and hostname.
            - Keep backups and restrict repository access if it contains sensitive data.
            """)
        else:
            st.markdown("""
            - Actualiza Python y las dependencias desde un entorno virtual x64 (Windows).
            - No compartas `.certs`, claves privadas, `deploy/.env` ni credenciales.
            - En producción publica solo 80/443 mediante Caddy; no abras 8501 a Internet.
            - Comprueba HTTPS desde el navegador y que la fecha/hostname coincida con el certificado.
            - Conserva copia de seguridad y restringe el acceso al repositorio si contiene datos sensibles.
            """)

elif vista == "Mejoras futuras":
    encabezado("Próximas mejoras", "Selecciona prioridades para orientar la siguiente versión; esta selección se procesa localmente y no se envía a ningún servidor.")
    opciones = [
        "Multiplicadores independientes por grupo",
        "Guardar escenarios con nombre y recuperarlos de forma segura",
        "Cálculo del impacto presupuestario por centro/campamento",
        "Importar y validar nuevas versiones de la plantilla .xls/.xlsx",
        "Exportación PDF con logo, firmas y campos de aprobación",
        "Control de usuarios, roles y bitácora de cambios",
        "Panel Windows con botones de iniciar/detener/abrir HTTPS",
        "Programar copias de seguridad y actualización de la escala",
        "Indicadores de nómina a partir de plazas autorizadas y verificadas",
        "Mejoras de accesibilidad y experiencia móvil (WCAG)",
        "Firma digital verificable de releases e informes con claves protegidas",
    ]
    opciones_visibles = [T(item) for item in opciones]
    seleccionadas_visibles = st.multiselect(T("¿Qué mejoras priorizarías?"), opciones_visibles, placeholder=T("Selecciona una o varias"))
    prioridades = ["Baja", "Media", "Alta", "Urgente"]
    prioridad_visible = st.select_slider(T("Prioridad global"), options=[T(item) for item in prioridades], value=T("Media"))
    prioridad = next(item for item in prioridades if T(item) == prioridad_visible)
    comentarios = st.text_area(T("Otras ideas o requisitos"), placeholder=T("Describe los cambios que necesitas en futuras versiones…"))
    resumen = [T("SUGERENCIAS DE MEJORA · CAPITAL HUMANO"), "", f"{T('Prioridad')}: {T(prioridad)}", "", T("Mejoras seleccionadas:")]
    resumen += [f"- {item}" for item in seleccionadas_visibles] or [f"- {T('No se seleccionaron opciones.')}"]
    resumen += ["", T("Comentarios:"), comentarios.strip() or T("(sin comentarios)"), ""]
    st.download_button(T("Descargar lista de mejoras (.txt)"), "\n".join(resumen), "Mejoras_Capital_Humano.txt", "text/plain")
    st.markdown(f"#### {T('Ideas recomendadas por impacto')}")
    if UI_LANGUAGE == "en":
        st.markdown("1. **Save and restore scenarios** for approved decisions.\n2. **Real budget impact** only when authorized occupied/vacant-position data exists.\n3. **Individual roles and audit log** before collaboration on salary data.\n4. **Source traceability and approval** in printable reports.")
    else:
        st.markdown("1. **Guardar y recuperar escenarios** para comparar decisiones aprobadas.\n2. **Presupuesto real** solo cuando exista un registro autorizado de plazas ocupadas/vacantes.\n3. **Roles individuales y bitácora** antes de permitir colaboración con datos salariales.\n4. **Trazabilidad de fuentes y aprobación** en los informes imprimibles.")
    st.caption(T("Consulta criterios de aceptación y prioridades completas en la hoja de ruta de Capital Humano. No adjuntes datos personales ni salarios reales a propuestas públicas de GitHub."))

st.markdown("---")
st.markdown(f"<div class='ch-muted'>{T('Capital Humano · versión {version} · herramienta de modelación y apoyo. Validar normativa e importes antes de aplicar una propuesta salarial.', version=version_visible)}</div>", unsafe_allow_html=True)
